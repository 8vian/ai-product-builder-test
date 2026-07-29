from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from ai_product_builder.cli import build_parser
from ai_product_builder.phase_b.errors import InputValidationError
from ai_product_builder.phase_b.manual_finalization import (
    MANUAL_FINALIZATION_ARTIFACTS,
    finalize_saved_run,
)
from ai_product_builder.phase_b.providers.apify import ApifyInstagramProvider


SOURCE_RUN = Path("output/phase_b/live-20260728T211907Z")
REVIEW_FILE = Path(
    "data/reviews/phase_b/"
    "live-20260728T211907Z-manual-final.json"
)
FINAL_USERNAMES = (
    "verkhovskaya_style",
    "olganemka_stylist",
    "angelashegiryan",
)


def _tree_hashes(path: Path) -> dict[str, str]:
    return {
        item.relative_to(path).as_posix(): hashlib.sha256(
            item.read_bytes()
        ).hexdigest()
        for item in sorted(path.rglob("*"))
        if item.is_file() and not item.name.startswith("~$")
    }


@pytest.fixture(scope="module")
def finalized_run(
    tmp_path_factory: pytest.TempPathFactory,
) -> tuple[Path, dict, tuple[dict, ...]]:
    output = tmp_path_factory.mktemp("manual-final") / "run"
    before = _tree_hashes(SOURCE_RUN)

    def _network_forbidden(*_args, **_kwargs):
        raise AssertionError("provider construction or request is forbidden")

    with (
        patch.object(
            ApifyInstagramProvider,
            "__init__",
            _network_forbidden,
        ),
        patch.object(
            ApifyInstagramProvider,
            "discover",
            _network_forbidden,
        ),
        patch.object(
            ApifyInstagramProvider,
            "enrich",
            _network_forbidden,
        ),
    ):
        manifest, paths, selected = finalize_saved_run(
            SOURCE_RUN,
            output,
            review_path=REVIEW_FILE,
        )

    assert before == _tree_hashes(SOURCE_RUN)
    assert {path.name for path in paths} == set(
        MANUAL_FINALIZATION_ARTIFACTS
    )
    return output, manifest, selected


def test_manual_finalize_cli_is_a_separate_offline_command() -> None:
    args = build_parser().parse_args(
        [
            "phase-b",
            "manual-finalize",
            "--source-run",
            str(SOURCE_RUN),
            "--review-file",
            str(REVIEW_FILE),
            "--output-dir",
            "output/phase_b/example-manual-final",
        ]
    )

    assert args.phase_b_command == "manual-finalize"
    assert not hasattr(args, "config")
    assert not hasattr(args, "mode")


def test_final_shortlist_is_exactly_three_in_human_order(
    finalized_run: tuple[Path, dict, tuple[dict, ...]],
) -> None:
    output, _, selected = finalized_run
    payload = json.loads(
        (output / "new_creators.json").read_text(encoding="utf-8")
    )

    assert tuple(item["username"] for item in selected) == FINAL_USERNAMES
    assert tuple(item["username"] for item in payload) == FINAL_USERNAMES
    assert [
        item["manual_verification_status"] for item in payload
    ] == ["approved", "approved", "pending"]
    assert {item["outreach_status"] for item in payload} == {"not_sent"}

    with (output / "eligible_candidates.csv").open(
        encoding="utf-8-sig", newline=""
    ) as stream:
        rows = list(csv.DictReader(stream))
    assert tuple(row["username"] for row in rows) == FINAL_USERNAMES


def test_three_offers_use_human_greetings_and_saved_fashion_posts(
    finalized_run: tuple[Path, dict, tuple[dict, ...]],
) -> None:
    output, _, selected = finalized_run
    greetings = ("Катя, здравствуйте!", "Ольга, здравствуйте!", "Ангела, здравствуйте!")
    expected_urls = (
        "https://www.instagram.com/p/DZE_YfrDegn/",
        "https://www.instagram.com/p/DbSlmTzNZUg/",
        "https://www.instagram.com/p/DbTRVFYtEgD/",
    )

    for item, greeting, post_url in zip(
        selected, greetings, expected_urls, strict=True
    ):
        assert item["barter_offer"].startswith(greeting)
        assert (
            "товар из новой коллекции LD Latte в обмен на согласованный "
            "Reels или нативный обзор"
        ) in item["barter_offer"]
        assert item["recent_post_url"] == post_url

    evidence = json.loads(
        (output / "selected_post_evidence.json").read_text(
            encoding="utf-8"
        )
    )
    assert tuple(item["post_url"] for item in evidence) == expected_urls
    assert all(
        item["direct_clothing_or_fashion_evidence"] is True
        for item in evidence
    )
    assert "plus-size" in selected[0]["barter_offer"]
    assert "баз" in selected[1]["barter_offer"].casefold()
    assert "ткань" in selected[2]["barter_offer"].casefold()


def test_rejected_profiles_have_exact_reasons_and_no_offers(
    finalized_run: tuple[Path, dict, tuple[dict, ...]],
) -> None:
    output, _, _ = finalized_run
    with (output / "excluded_candidates.csv").open(
        encoding="utf-8-sig", newline=""
    ) as stream:
        exclusions = {
            row["username"]: row for row in csv.DictReader(stream)
        }
    assert exclusions["ayuma.style"]["exclusion_reason"] == (
        "commercial_conflict_own_fashion_brand"
    )
    assert exclusions["miss_sunrise9"]["exclusion_reason"] == (
        "commercial_conflict_clothing_showroom"
    )
    assert exclusions["ayuma.style"]["commercial_conflict_account"] == (
        "@dorogaya.brand"
    )
    assert exclusions["miss_sunrise9"]["commercial_conflict_account"] == (
        "@peperoncino_shop_italy_"
    )
    assert exclusions["stylistelenaialena"]["exclusion_reason"] == (
        "professional_portfolio_not_influencer_creator"
    )
    stylist_evidence = json.loads(
        exclusions["stylistelenaialena"]["manual_review_evidence"]
    )
    assert {item["signal_type"] for item in stylist_evidence} == {
        "professional_specialization",
        "portfolio_commercial_shoots",
        "regular_short_video_content",
        "influencer_format_gap",
    }
    assert {
        item["source_field"] for item in stylist_evidence
    } == {
        "profile.full_name",
        "profile.biography",
        "profile.recent_posts.caption",
        "candidate.metrics.short_video_share",
        "candidate.metrics.post_format_distribution",
    }

    audit = {
        item["username"]: item
        for item in json.loads(
            (output / "eligible_audit_pool.json").read_text(
                encoding="utf-8"
            )
        )
    }
    for username in (
        "ayuma.style",
        "miss_sunrise9",
        "stylistelenaialena",
    ):
        assert audit[username]["status"] == "rejected"
        assert audit[username]["barter_offer"] == ""
        assert audit[username]["outreach_status"] == "not_sent"


def test_professional_portfolio_is_rejected_without_offer(
    finalized_run: tuple[Path, dict, tuple[dict, ...]],
) -> None:
    output, _, _ = finalized_run
    audit = {
        item["username"]: item
        for item in json.loads(
            (output / "eligible_audit_pool.json").read_text(
                encoding="utf-8"
            )
        )
    }
    item = audit["stylistelenaialena"]

    assert item["status"] == "rejected"
    assert item["manual_review_reason"] == (
        "professional_portfolio_not_influencer_creator"
    )
    assert item["eligibility_status"] == "ineligible"
    assert item["campaign_bucket"] == "ineligible_or_insufficient"
    assert item["eligibility_reasons"][-1] == (
        "professional_portfolio_not_influencer_creator"
    )
    assert item["barter_offer"] == ""
    assert item["manual_verification_status"] == "rejected"
    assert item["outreach_status"] == "not_sent"
    assert len(item["manual_review_evidence"]) == 5


def test_pending_manager_note_is_exact_and_draft_is_not_sent(
    finalized_run: tuple[Path, dict, tuple[dict, ...]],
) -> None:
    output, _, selected = finalized_run
    angela = selected[2]
    exact_note = (
        "Требуется согласовать с менеджером возможность бартера, формат "
        "интеграции и состав предоставляемого контента"
    )
    drafts = (output / "barter_offer_drafts.md").read_text(
        encoding="utf-8"
    )

    assert angela["verification_notes"] == exact_note
    assert angela["manual_verification_status"] == "pending"
    assert angela["outreach_status"] == "not_sent"
    assert "preliminary only" in drafts
    assert drafts.count("## ") == 3
    assert "ayuma.style" not in drafts
    assert "miss_sunrise9" not in drafts
    assert "stylistelenaialena" not in drafts


def test_report_contains_all_human_exclusions_and_omission_reasons(
    finalized_run: tuple[Path, dict, tuple[dict, ...]],
) -> None:
    output, _, _ = finalized_run
    report = (output / "discovery_report.md").read_text(
        encoding="utf-8"
    )

    assert "## Human-in-the-loop exclusions" in report
    assert "commercial_conflict_own_fashion_brand" in report
    assert "commercial_conflict_clothing_showroom" in report
    assert "professional_portfolio_not_influencer_creator" in report
    assert "2 short-video posts out of 12 sampled posts" in report
    assert "All three manually rejected profiles have no offer." in report
    assert "Provider requests made: `0`" in report
    assert "Budget spent: `$0.00`" in report
    assert "Outreach messages sent: `0`" in report


def test_manifest_records_zero_requests_cost_and_messages(
    finalized_run: tuple[Path, dict, tuple[dict, ...]],
) -> None:
    output, manifest, _ = finalized_run
    disk_manifest = json.loads(
        (output / "run_manifest.json").read_text(encoding="utf-8")
    )

    assert disk_manifest == manifest
    assert manifest["offline_reselection"] is True
    assert manifest["manual_finalization"] is True
    assert manifest["human_in_the_loop_review"] is True
    assert manifest["provider_requests_made"] == 0
    assert manifest["budget_spent_usd"] == 0.0
    assert manifest["outreach_messages_sent"] == 0
    assert manifest["source_run_unchanged"] is True
    assert manifest["phase_a_unchanged"] is True
    assert manifest["source_run_path"] == SOURCE_RUN.as_posix()
    assert manifest["artifacts"] == {
        name: name for name in MANUAL_FINALIZATION_ARTIFACTS
    }
    assert all(
        not Path(value).is_absolute()
        for value in manifest["artifacts"].values()
    )
    assert manifest["counts"] == {
        "source_eligible_candidates": 6,
        "manual_reviewed_candidates": 6,
        "final_selected_candidates": 3,
        "preserved_automatic_exclusions": 32,
        "manual_commercial_conflict_exclusions": 2,
        "manual_other_exclusions": 1,
        "excluded_records_total": 35,
        "manual_approved": 2,
        "manual_pending": 1,
        "manual_rejected": 3,
        "eligible_not_selected": 0,
    }


def test_workbook_preserves_source_sheet_and_has_exact_rows_and_links(
    finalized_run: tuple[Path, dict, tuple[dict, ...]],
) -> None:
    output, _, _ = finalized_run
    workbook = load_workbook(
        output / "Блогеры_phase_b.xlsx",
        read_only=False,
        data_only=False,
    )
    try:
        assert workbook.sheetnames == ["Исходник", "Новые блоггеры"]
        worksheet = workbook["Новые блоггеры"]
        headers = {
            cell.value: cell.column for cell in worksheet[1]
        }
        usernames = [
            worksheet.cell(row, headers["username"]).value
            for row in range(2, worksheet.max_row + 1)
        ]
        assert tuple(usernames) == FINAL_USERNAMES
        assert worksheet.max_row == 4
        for row in range(2, 5):
            profile = worksheet.cell(row, headers["profile_url"])
            post = worksheet.cell(row, headers["recent_post_url"])
            assert profile.hyperlink is not None
            assert profile.hyperlink.target == profile.value
            assert post.hyperlink is not None
            assert post.hyperlink.target == post.value
            assert (
                worksheet.cell(row, headers["outreach_status"]).value
                == "not_sent"
            )
        assert (
            worksheet.cell(
                2, headers["manual_verification_status"]
            ).value
            == "approved"
        )
        assert (
            worksheet.cell(
                4, headers["manual_verification_status"]
            ).value
            == "pending"
        )
    finally:
        workbook.close()


def test_review_must_cover_the_complete_saved_eligible_pool(
    tmp_path: Path,
) -> None:
    payload = json.loads(REVIEW_FILE.read_text(encoding="utf-8"))
    payload["decisions"] = payload["decisions"][:-1]
    review = tmp_path / "incomplete.json"
    review.write_text(
        json.dumps(payload, ensure_ascii=False),
        encoding="utf-8",
    )

    with pytest.raises(InputValidationError, match="must be reviewed"):
        finalize_saved_run(
            SOURCE_RUN,
            tmp_path / "output",
            review_path=review,
        )


def test_duplicate_review_decision_is_rejected(tmp_path: Path) -> None:
    payload = json.loads(REVIEW_FILE.read_text(encoding="utf-8"))
    payload["decisions"].append(dict(payload["decisions"][0]))
    review = tmp_path / "duplicate.json"
    review.write_text(
        json.dumps(payload, ensure_ascii=False),
        encoding="utf-8",
    )

    with pytest.raises(InputValidationError, match="duplicate decision"):
        finalize_saved_run(
            SOURCE_RUN,
            tmp_path / "output",
            review_path=review,
        )


def test_lifestyle_post_cannot_ground_a_clothing_offer(
    tmp_path: Path,
) -> None:
    payload = json.loads(REVIEW_FILE.read_text(encoding="utf-8"))
    decision = next(
        item
        for item in payload["decisions"]
        if item["username"] == "verkhovskaya_style"
    )
    decision["selected_post_url"] = (
        "https://www.instagram.com/p/DZnhAXPmS6H/"
    )
    review = tmp_path / "lifestyle.json"
    review.write_text(
        json.dumps(payload, ensure_ascii=False),
        encoding="utf-8",
    )

    with pytest.raises(
        InputValidationError,
        match="lacks direct clothing/fashion evidence",
    ):
        finalize_saved_run(
            SOURCE_RUN,
            tmp_path / "output",
            review_path=review,
        )


def test_manual_exclusion_evidence_must_match_saved_source(
    tmp_path: Path,
) -> None:
    payload = json.loads(REVIEW_FILE.read_text(encoding="utf-8"))
    decision = next(
        item
        for item in payload["decisions"]
        if item["username"] == "stylistelenaialena"
    )
    decision["manual_exclusion_evidence"][0][
        "evidence_text"
    ] = "Invented professional title"
    review = tmp_path / "invented-evidence.json"
    review.write_text(
        json.dumps(payload, ensure_ascii=False),
        encoding="utf-8",
    )

    with pytest.raises(
        InputValidationError,
        match="evidence is absent from saved profile.full_name",
    ):
        finalize_saved_run(
            SOURCE_RUN,
            tmp_path / "output",
            review_path=review,
        )


def test_derived_manual_exclusion_evidence_must_match_saved_metrics(
    tmp_path: Path,
) -> None:
    payload = json.loads(REVIEW_FILE.read_text(encoding="utf-8"))
    decision = next(
        item
        for item in payload["decisions"]
        if item["username"] == "stylistelenaialena"
    )
    derived = next(
        item
        for item in decision["manual_exclusion_evidence"]
        if item["source_field"] == "candidate.metrics.short_video_share"
    )
    derived["evidence_text"] = (
        "Observed 12 short-video posts out of 12 sampled posts "
        "(share 1.000000)."
    )
    review = tmp_path / "invented-metrics.json"
    review.write_text(
        json.dumps(payload, ensure_ascii=False),
        encoding="utf-8",
    )

    with pytest.raises(
        InputValidationError,
        match="does not match saved post-format data",
    ):
        finalize_saved_run(
            SOURCE_RUN,
            tmp_path / "output",
            review_path=review,
        )


def test_rejected_decision_requires_rejected_manual_status(
    tmp_path: Path,
) -> None:
    payload = json.loads(REVIEW_FILE.read_text(encoding="utf-8"))
    decision = next(
        item
        for item in payload["decisions"]
        if item["username"] == "stylistelenaialena"
    )
    decision["manual_verification_status"] = "pending"
    review = tmp_path / "pending-rejection.json"
    review.write_text(
        json.dumps(payload, ensure_ascii=False),
        encoding="utf-8",
    )

    with pytest.raises(
        InputValidationError,
        match="must have manual_verification_status=rejected",
    ):
        finalize_saved_run(
            SOURCE_RUN,
            tmp_path / "output",
            review_path=review,
        )
