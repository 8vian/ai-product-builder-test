from __future__ import annotations

import csv
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from ai_product_builder.phase_b.io import PHASE_B_COLUMNS
from ai_product_builder.phase_b.io.google_sheets import GoogleSheetsAdapter
from ai_product_builder.phase_b.io.local_csv import (
    write_candidate_csv,
    write_excluded_candidates_csv,
)
from ai_product_builder.phase_b.io.local_xlsx import write_candidates_workbook


SHEET_NAME = "Новые блоггеры"


def _candidate(
    username: str = "new.creator__",
    *,
    status: str = "pending",
    notes: str = "",
    score: float = 81.25,
) -> dict[str, object]:
    return {
        "platform": "instagram",
        "username": username,
        "profile_url": f"https://www.instagram.com/{username}/",
        "followers": 12_500,
        "median_likes": 220.5,
        "median_comments": 14.0,
        "engagement_rate": 1.876,
        "usable_posts": 8,
        "sampled_posts": 10,
        "data_completeness": 0.8,
        "short_video_share": 0.6,
        "last_post_date": "2026-01-28T00:00:00+00:00",
        "score": score,
        "score_components": [
            {
                "name": "content_and_aesthetic_fit",
                "score": 24.0,
                "max_score": 30.0,
            }
        ],
        "selection_explanation": "Ranked from the complete eligible pool.",
        "evidence": [
            {
                "signal_type": "fashion",
                "source_field": "recent_posts.caption",
                "source_reference": "post-1",
                "evidence_text": "Observed fashion term",
                "observation_type": "direct",
                "url": "https://www.instagram.com/p/newcreator1/",
            }
        ],
        "recent_post_url": "https://www.instagram.com/p/newcreator1/",
        "barter_offer": "Draft requiring manual approval.",
        "manual_verification_status": status,
        "verification_notes": notes,
        "discovery_confidence": 0.91,
        "eligibility_status": "eligible",
        "eligibility_reasons": [],
        "query_ids": ["fashion_style_01", "reels_integrations_01"],
        "provider": "fixture",
        "collected_at": "2026-02-01T00:00:00+00:00",
        "offer_generation_mode": "deterministic_template",
        "source_exclusion_check": "clear",
    }


def _source_workbook(path: Path) -> None:
    workbook = Workbook()
    original = workbook.active
    original.title = "Блогеры"
    original["A1"] = "source identity"
    original["A2"] = "_crazy___unicorn_"
    original["A2"].font = Font(bold=True, color="CC0000")
    original["A2"].fill = PatternFill("solid", fgColor="FFF2CC")
    untouched = workbook.create_sheet("Unrelated")
    untouched["C7"] = "=1+2"
    untouched["D8"] = "keep me"
    workbook.save(path)
    workbook.close()


def _file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_candidate_csv_has_bom_stable_headers_and_json_nested_fields(
    tmp_path: Path,
) -> None:
    output = tmp_path / "new_creators.csv"
    write_candidate_csv(output, [_candidate()])

    assert output.read_bytes().startswith(b"\xef\xbb\xbf")
    with output.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert tuple(rows[0]) == PHASE_B_COLUMNS
    assert rows[0]["username"] == "new.creator__"
    assert json.loads(rows[0]["score_components"])[0]["score"] == 24.0
    assert json.loads(rows[0]["evidence"])[0]["source_reference"] == "post-1"
    assert json.loads(rows[0]["query_ids"]) == [
        "fashion_style_01",
        "reels_integrations_01",
    ]


def test_excluded_csv_nested_fields_are_single_json_values(
    tmp_path: Path,
) -> None:
    output = tmp_path / "excluded_candidates.csv"
    write_excluded_candidates_csv(
        output,
        [
            {
                "platform": "instagram",
                "username": "__aparina",
                "profile_url": "https://www.instagram.com/__aparina/",
                "normalized_username": "__aparina",
                "exclusion_reason": "source_exclusion",
                "exclusion_reasons": ["historical_alias"],
                "query_ids": ["fashion_style_01"],
                "provider": "fixture",
                "exclusion_registry_match": {
                    "normalized_username": "__aparina",
                    "reasons": ["historical_alias"],
                },
            }
        ],
    )

    with output.open(encoding="utf-8-sig", newline="") as handle:
        row = next(csv.DictReader(handle))

    assert json.loads(row["exclusion_reasons"]) == ["historical_alias"]
    assert json.loads(row["query_ids"]) == ["fashion_style_01"]
    assert json.loads(row["exclusion_registry_match"]) == {
        "normalized_username": "__aparina",
        "reasons": ["historical_alias"],
    }


def test_xlsx_preserves_all_source_sheets_and_never_overwrites_raw_input(
    tmp_path: Path,
) -> None:
    source = tmp_path / "Блогеры.xlsx"
    output = tmp_path / "Блогеры_phase_b.xlsx"
    _source_workbook(source)
    before = _file_hash(source)

    write_candidates_workbook(
        source, output, [_candidate()], sheet_name=SHEET_NAME
    )

    assert source.exists()
    assert _file_hash(source) == before
    assert output.exists()
    workbook = load_workbook(output, data_only=False)
    try:
        assert workbook.sheetnames == ["Блогеры", "Unrelated", SHEET_NAME]
        assert workbook["Блогеры"]["A1"].value == "source identity"
        assert workbook["Блогеры"]["A2"].value == "_crazy___unicorn_"
        assert workbook["Блогеры"]["A2"].font.bold is True
        assert workbook["Блогеры"]["A2"].fill.fgColor.rgb.endswith("FFF2CC")
        assert workbook["Unrelated"]["C7"].value == "=1+2"
        assert workbook["Unrelated"]["D8"].value == "keep me"
    finally:
        workbook.close()


def test_xlsx_uses_exact_sheet_headers_and_real_hyperlinks(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "result.xlsx"
    _source_workbook(source)
    candidate = _candidate("dots._and__underscores")

    write_candidates_workbook(
        source, output, [candidate], sheet_name=SHEET_NAME
    )

    workbook = load_workbook(output, data_only=False)
    try:
        worksheet = workbook[SHEET_NAME]
        assert worksheet.title == "Новые блоггеры"
        assert tuple(cell.value for cell in worksheet[1]) == PHASE_B_COLUMNS
        columns = {
            cell.value: cell.column for cell in worksheet[1] if cell.value
        }
        assert worksheet.cell(2, columns["username"]).value == (
            "dots._and__underscores"
        )
        profile = worksheet.cell(2, columns["profile_url"])
        post = worksheet.cell(2, columns["recent_post_url"])
        assert profile.hyperlink is not None
        assert profile.hyperlink.target == candidate["profile_url"]
        assert post.hyperlink is not None
        assert post.hyperlink.target == candidate["recent_post_url"]
        assert worksheet.freeze_panes == "A2"
        assert worksheet.data_validations.count == 1
    finally:
        workbook.close()


def test_xlsx_rerun_preserves_manual_status_and_notes_by_normalized_username(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "result.xlsx"
    _source_workbook(source)
    write_candidates_workbook(
        source, output, [_candidate("Mixed._Case")], sheet_name=SHEET_NAME
    )

    workbook = load_workbook(output)
    worksheet = workbook[SHEET_NAME]
    columns = {cell.value: cell.column for cell in worksheet[1] if cell.value}
    worksheet.cell(2, columns["manual_verification_status"]).value = "approved"
    worksheet.cell(2, columns["verification_notes"]).value = "Pavel checked link"
    workbook.save(output)
    workbook.close()

    rerun = _candidate("mixed._case", status="pending", notes="", score=88.0)
    write_candidates_workbook(source, output, [rerun], sheet_name=SHEET_NAME)

    workbook = load_workbook(output, data_only=False)
    try:
        worksheet = workbook[SHEET_NAME]
        columns = {cell.value: cell.column for cell in worksheet[1] if cell.value}
        assert worksheet.cell(2, columns["score"]).value == 88.0
        assert (
            worksheet.cell(2, columns["manual_verification_status"]).value
            == "approved"
        )
        assert (
            worksheet.cell(2, columns["verification_notes"]).value
            == "Pavel checked link"
        )
    finally:
        workbook.close()


def test_xlsx_refuses_to_use_raw_workbook_as_output(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    _source_workbook(source)
    before = _file_hash(source)

    try:
        write_candidates_workbook(source, source, [_candidate()])
    except ValueError as exc:
        assert "must not overwrite" in str(exc)
    else:
        raise AssertionError("writer allowed raw workbook overwrite")
    assert _file_hash(source) == before


class _FakeSheetsClient:
    def __init__(self, values: list[list[object]]) -> None:
        self.values = values
        self.reads: list[tuple[str, str]] = []
        self.batches: list[tuple[str, list[dict[str, object]]]] = []

    def get_values(self, spreadsheet_id: str, range_name: str) -> dict[str, object]:
        self.reads.append((spreadsheet_id, range_name))
        return {"values": self.values}

    def batch_update_values(
        self, spreadsheet_id: str, updates: list[dict[str, object]]
    ) -> dict[str, int]:
        self.batches.append((spreadsheet_id, updates))
        return {"updated": len(updates)}


def test_google_sheets_upsert_batches_changes_and_preserves_manual_columns() -> None:
    existing_row = ["" for _ in PHASE_B_COLUMNS]
    existing_row[PHASE_B_COLUMNS.index("username")] = "Mixed._Case"
    existing_row[PHASE_B_COLUMNS.index("profile_url")] = (
        "https://www.instagram.com/Mixed._Case/"
    )
    existing_row[
        PHASE_B_COLUMNS.index("manual_verification_status")
    ] = "needs_review"
    existing_row[PHASE_B_COLUMNS.index("verification_notes")] = "check identity"
    client = _FakeSheetsClient([list(PHASE_B_COLUMNS), existing_row])
    adapter = GoogleSheetsAdapter(
        client, "spreadsheet-1", worksheet_name=SHEET_NAME
    )

    receipt = adapter.upsert(
        [_candidate("mixed._case"), _candidate("another.creator")]
    )

    assert receipt.inserted_rows == 1
    assert receipt.updated_rows == 1
    assert receipt.preserved_manual_rows == 1
    assert len(client.batches) == 1
    _, updates = client.batches[0]
    assert len(updates) == 3
    assert updates[0]["range"] == "'Новые блоггеры'!A1:AB1"
    assert all(
        str(update["range"]).startswith("'Новые блоггеры'!")
        for update in updates
    )
    assert not any("Unrelated" in str(update["range"]) for update in updates)

    updated_existing = updates[1]["values"][0]  # type: ignore[index]
    assert (
        updated_existing[
            PHASE_B_COLUMNS.index("manual_verification_status")
        ]
        == "needs_review"
    )
    assert (
        updated_existing[PHASE_B_COLUMNS.index("verification_notes")]
        == "check identity"
    )


def test_google_sheets_adapter_requires_injected_credentials_client() -> None:
    adapter = GoogleSheetsAdapter(None, "spreadsheet-1")
    assert adapter.configured is False
    try:
        adapter.upsert([_candidate()])
    except RuntimeError as exc:
        assert "not configured" in str(exc)
    else:
        raise AssertionError("unconfigured Google Sheets adapter performed a write")
