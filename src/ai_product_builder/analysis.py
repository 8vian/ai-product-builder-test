from __future__ import annotations

import json
import math
import os
import re
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from statistics import median
from typing import Any, Iterable
from urllib.parse import urlparse

from openpyxl import load_workbook

from .models import (
    PostFormat,
    Profile,
    ProfileMetrics,
    ProfileScore,
    ProfileStatus,
    ScoreComponent,
    SignalEvidence,
    SignalResult,
)

BRAND_USERNAMES = {"nike", "apple"}
MIN_SAMPLED_POSTS = 3
MIN_USABLE_ENGAGEMENT_POSTS = 6

SIGNAL_TERMS: dict[str, tuple[str, ...]] = {
    "commercial_pr": (
        "реклама",
        "сотруднич",
        "сотрудничество",
        "collab",
        "commercial",
        "амбассадор",
        "ambassador",
        "partnership",
        "по рекламе",
        "brand work",
        "brand collaboration",
        "работа с брендами",
        "сотрудничество с брендами",
        "бартер",
        "barter",
    ),
    "contact": (
        "direct",
        "директ",
        "manager",
        "менеджер",
        "почта",
        "email",
        "e-mail",
        "telegram",
        "телеграм",
        "tg:",
    ),
    "fashion": (
        "fashion",
        "мода",
        "стиль",
        "одежд",
        "лук",
        "outfit",
        "гардероб",
        "женствен",
    ),
    "beauty": (
        "beauty",
        "бьюти",
        "космет",
        "макияж",
        "уход",
        "визаж",
        "skin",
        "hair",
    ),
    "lifestyle": (
        "lifestyle",
        "life",
        "лайф",
        "жизн",
        "дом",
        "семья",
        "мама",
        "путеше",
        "вдохнов",
    ),
    "ugc": ("ugc", "content creator", "контент креатор", "creator"),
    "marketplace": (
        "wildberries",
        "вайлдберриз",
        "wb",
        "ozon",
        "маркетплейс",
        "распаков",
        "обзор",
        "покупок",
        "находк",
    ),
}

EMAIL_PATTERN = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+", re.IGNORECASE)
PR_PATTERN = re.compile(r"\bpr\b", re.IGNORECASE)
INSTAGRAM_USERNAME_PATTERN = re.compile(r"^[A-Za-z0-9._]+$")
NATIVE_INTEGRATION_TERMS = (
    "review",
    "обзор",
    "unbox",
    "распаков",
    "try-on",
    "try on",
    "примерк",
    "promo code",
    "promocode",
    "промокод",
    "артикул",
    "product",
    "товар",
    "покупк",
    "ссылка",
    "http://",
    "https://",
)
CONTACT_URL_MARKERS = (
    "t.me/",
    "telegram",
    "wa.me/",
    "whatsapp",
    "mailto:",
    "vk.com/",
    "vk.ru/",
)


@dataclass(slots=True)
class WorkbookReference:
    row: int
    ordinal: int | None
    display_value: str
    hyperlink_target: str | None
    selected_url: str
    selected_username: str | None
    display_username: str | None
    hyperlink_preferred: bool
    display_target_mismatch: bool

    def serializable(self) -> dict[str, Any]:
        return {
            "row": self.row,
            "ordinal": self.ordinal,
            "display_value": self.display_value,
            "hyperlink_target": self.hyperlink_target,
            "selected_url": self.selected_url,
            "selected_username": self.selected_username,
            "display_username": self.display_username,
            "hyperlink_preferred": self.hyperlink_preferred,
            "display_target_mismatch": self.display_target_mismatch,
        }


@dataclass(slots=True)
class AnalysisResult:
    generated_at: datetime
    analysis_as_of: datetime
    profiles: list[Profile]
    metrics: dict[int, ProfileMetrics]
    scores: dict[int, ProfileScore]
    workbook_references: list[WorkbookReference]
    workbook_reconciliation: dict[str, Any]
    manual_audit: dict[str, Any]
    dataset_summary: dict[str, Any]
    brand_reference_signals: list[dict[str, Any]]
    ideal_creator_profile: dict[str, Any]
    validation_issues: list[str]


def read_json(path: Path) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ValueError(f"{path}: invalid JSON at line {exc.lineno}, column {exc.colno}") from exc


def normalize_username_from_url(value: str | None) -> str | None:
    if not value:
        return None
    text = value.strip()
    match = re.search(r"@([A-Za-z0-9._]+)", text)
    if match and "instagram" not in text.casefold():
        return match.group(1).casefold()
    if not re.match(r"^https?://", text, re.IGNORECASE):
        candidate = text.lstrip("@").split("?", 1)[0].strip("/")
        return candidate.casefold() if INSTAGRAM_USERNAME_PATTERN.match(candidate) else None
    parsed = urlparse(text)
    if "instagram.com" not in parsed.netloc.casefold():
        return None
    parts = [part for part in parsed.path.split("/") if part]
    if not parts:
        return None
    candidate = parts[0].lstrip("@")
    if candidate.casefold() in {"p", "reel", "reels", "stories", "explore"}:
        return None
    return candidate.casefold() if INSTAGRAM_USERNAME_PATTERN.match(candidate) else None


def read_workbook_references(path: Path) -> list[WorkbookReference]:
    workbook = load_workbook(path, read_only=False, data_only=False)
    sheet = workbook["Исходник"] if "Исходник" in workbook.sheetnames else workbook.worksheets[0]
    references: list[WorkbookReference] = []
    for row in range(1, sheet.max_row + 1):
        ordinal_cell = sheet.cell(row, 1)
        link_cell = sheet.cell(row, 2)
        display = str(link_cell.value or "").strip()
        target = link_cell.hyperlink.target if link_cell.hyperlink else None
        if not display and not target:
            continue
        selected = str(target or display).strip()
        display_username = normalize_username_from_url(display)
        target_username = normalize_username_from_url(target)
        selected_username = target_username or display_username
        mismatch = bool(
            target
            and display
            and (
                display != target
                or (
                    display_username
                    and target_username
                    and display_username != target_username
                )
            )
        )
        try:
            ordinal = int(ordinal_cell.value) if ordinal_cell.value is not None else None
        except (TypeError, ValueError):
            ordinal = None
        references.append(
            WorkbookReference(
                row=row,
                ordinal=ordinal,
                display_value=display,
                hyperlink_target=target,
                selected_url=selected,
                selected_username=selected_username,
                display_username=display_username,
                hyperlink_preferred=bool(target),
                display_target_mismatch=mismatch,
            )
        )
    workbook.close()
    return references


def classify_profile(profile: Profile) -> None:
    username = (profile.username or normalize_username_from_url(profile.input_url) or "").casefold()
    usable_engagement_posts = sum(
        post.likes is not None and post.comments is not None
        for post in profile.latest_posts
    )
    if profile.error and profile.error.casefold() in {"not_found", "unresolved"}:
        profile.status = ProfileStatus.UNRESOLVED_NOT_FOUND
        profile.exclusion_reason = (
            f"Source status is {profile.error}; missing values were preserved."
        )
    elif profile.private is True:
        profile.status = ProfileStatus.PRIVATE
        profile.exclusion_reason = (
            "Private profile: latest-post metrics are unavailable and were not inferred."
        )
    elif username in BRAND_USERNAMES:
        profile.status = ProfileStatus.BRAND_REFERENCE
        profile.exclusion_reason = (
            "Nike/Apple brand reference: used only for content-format alignment signals."
        )
    elif (
        not profile.username
        or profile.followers is None
        or profile.followers <= 0
        or len(profile.latest_posts) < MIN_SAMPLED_POSTS
        or usable_engagement_posts < MIN_USABLE_ENGAGEMENT_POSTS
    ):
        profile.status = ProfileStatus.INSUFFICIENT_DATA
        profile.exclusion_reason = (
            "Insufficient public data for robust quantitative scoring "
            f"(requires username, positive followers, {MIN_SAMPLED_POSTS}+ sampled "
            f"posts, and {MIN_USABLE_ENGAGEMENT_POSTS}+ posts with usable likes and comments; "
            f"found {usable_engagement_posts} usable engagement posts)."
        )
    else:
        profile.status = ProfileStatus.CREATOR
        profile.exclusion_reason = None


def normalized_text(profile: Profile) -> str:
    values = [profile.full_name, profile.biography]
    for post in profile.latest_posts:
        values.append(post.caption)
        values.extend(post.hashtags)
    return "\n".join(values).casefold()


def detect_signals(profile: Profile) -> dict[str, SignalResult]:
    text_sources: list[tuple[str, str, str]] = [
        ("fullName", "profile", profile.full_name),
        ("biography", "profile", profile.biography),
    ]
    for post_index, post in enumerate(profile.latest_posts):
        reference = post.post_id or f"index:{post_index}"
        text_sources.append(("latestPosts.caption", reference, post.caption))
        if post.hashtags:
            text_sources.append(
                ("latestPosts.hashtags", reference, " ".join(post.hashtags))
            )

    results: dict[str, SignalResult] = {}
    for name, terms in SIGNAL_TERMS.items():
        matches: set[str] = set()
        provenance: list[SignalEvidence] = []
        for source_field, source_reference, source_text in text_sources:
            folded = source_text.casefold()
            for term in terms:
                if term.casefold() not in folded:
                    continue
                normalized_term = term.strip()
                matches.add(normalized_term)
                provenance.append(
                    SignalEvidence(
                        signal_type=name,
                        source_field=source_field,
                        source_reference=source_reference,
                        evidence_text=f"Matched explicit term: {normalized_term}",
                    )
                )
            if name == "commercial_pr" and PR_PATTERN.search(source_text):
                matches.add("pr")
                provenance.append(
                    SignalEvidence(
                        signal_type=name,
                        source_field=source_field,
                        source_reference=source_reference,
                        evidence_text="Matched explicit standalone term: PR",
                    )
                )
        ordered_matches = tuple(sorted(matches))
        evidence = (
            (f"Matched terms: {', '.join(ordered_matches[:8])}",)
            if ordered_matches
            else ()
        )
        results[name] = SignalResult(
            bool(ordered_matches),
            ordered_matches,
            evidence,
            tuple(provenance),
        )

    email_matches = tuple(sorted(set(EMAIL_PATTERN.findall(profile.biography))))
    contact_evidence = list(results["contact"].evidence)
    contact_provenance = list(results["contact"].provenance)
    if email_matches:
        contact_evidence.append(f"Email pattern in bio: {', '.join(email_matches)}")
        contact_provenance.extend(
            SignalEvidence(
                signal_type="contact",
                source_field="biography",
                source_reference="profile",
                evidence_text=f"Observed email address: {email}",
            )
            for email in email_matches
        )
    contact_urls = tuple(
        url
        for url in profile.external_urls
        if any(marker in url.casefold() for marker in CONTACT_URL_MARKERS)
    )
    if contact_urls:
        contact_evidence.append(f"{len(contact_urls)} contact-capable external URL(s)")
        contact_provenance.extend(
            SignalEvidence(
                signal_type="contact",
                source_field="externalUrls",
                source_reference="profile",
                evidence_text=f"Observed external contact/link URL: {url}",
            )
            for url in contact_urls
        )
    results["contact"] = SignalResult(
        results["contact"].detected or bool(email_matches) or bool(contact_urls),
        tuple(sorted(set(results["contact"].matches + email_matches))),
        tuple(contact_evidence),
        tuple(contact_provenance),
    )

    commercial_provenance = list(results["commercial_pr"].provenance)
    commercial_evidence = list(results["commercial_pr"].evidence)
    commercial_matches = set(results["commercial_pr"].matches)
    paid_partnership_posts = sum(post.paid_partnership for post in profile.latest_posts)
    for post_index, post in enumerate(profile.latest_posts):
        if not post.paid_partnership:
            continue
        reference = post.post_id or f"index:{post_index}"
        commercial_matches.add("paid-partnership-flag")
        commercial_provenance.append(
            SignalEvidence(
                signal_type="commercial_pr",
                source_field="latestPosts.paidPartnership",
                source_reference=reference,
                evidence_text="Observed paid-partnership flag.",
            )
        )
    if paid_partnership_posts:
        commercial_evidence.append(
            f"{paid_partnership_posts} sampled post(s) have a paid-partnership flag."
        )
    results["commercial_pr"] = SignalResult(
        bool(commercial_matches),
        tuple(sorted(set(commercial_matches))),
        tuple(dict.fromkeys(commercial_evidence)),
        tuple(commercial_provenance),
    )

    integration_posts = 0
    integration_examples: list[str] = []
    integration_provenance: list[SignalEvidence] = []
    integration_terms = tuple(
        dict.fromkeys((*SIGNAL_TERMS["marketplace"], *NATIVE_INTEGRATION_TERMS))
    )
    for post_index, post in enumerate(profile.latest_posts):
        reference = post.post_id or f"index:{post_index}"
        post_observations: list[SignalEvidence] = []
        if post.mentions:
            post_observations.append(
                SignalEvidence(
                    signal_type="native_product_integration",
                    source_field="latestPosts.mentions",
                    source_reference=reference,
                    evidence_text=(
                        "Observed structured account mention(s): "
                        + ", ".join(post.mentions)
                    ),
                )
            )
        if post.tagged_usernames:
            post_observations.append(
                SignalEvidence(
                    signal_type="native_product_integration",
                    source_field="latestPosts.taggedUsers",
                    source_reference=reference,
                    evidence_text=(
                        "Observed tagged account(s): "
                        + ", ".join(post.tagged_usernames)
                    ),
                )
            )
        caption_matches = sorted(
            {
                term.strip()
                for term in integration_terms
                if term.casefold() in post.caption.casefold()
            }
        )
        if caption_matches:
            post_observations.append(
                SignalEvidence(
                    signal_type="native_product_integration",
                    source_field="latestPosts.caption",
                    source_reference=reference,
                    evidence_text=(
                        "Observed product-focused caption term(s): "
                        + ", ".join(caption_matches[:8])
                    ),
                )
            )
        if post_observations:
            integration_posts += 1
            integration_provenance.extend(post_observations)
            if len(integration_examples) < 3 and post.post_id:
                integration_examples.append(post.post_id)
    results["native_product_integration"] = SignalResult(
        integration_posts > 0,
        (str(integration_posts),) if integration_posts else (),
        (
            f"{integration_posts}/{len(profile.latest_posts)} sampled posts contain "
            "mentions, tags, product-focused captions, links, promo codes, reviews, "
            "try-ons, unboxings, or marketplace/product language.",
            *(
                (f"Example post IDs: {', '.join(integration_examples)}",)
                if integration_examples
                else ()
            ),
        ),
        tuple(integration_provenance),
    )
    return results


def safe_median(values: Iterable[float | int | None]) -> float | None:
    cleaned = [float(value) for value in values if value is not None]
    return median(cleaned) if cleaned else None


def calculate_metrics(profile: Profile, as_of: datetime) -> ProfileMetrics:
    posts = list(profile.latest_posts)
    likes = [post.likes for post in posts if post.likes is not None]
    comments = [post.comments for post in posts if post.comments is not None]
    engagement_rates = [
        ((post.likes + post.comments) / profile.followers) * 100
        for post in posts
        if profile.followers
        and profile.followers > 0
        and post.likes is not None
        and post.comments is not None
    ]
    format_counts_counter = Counter(post.format.value for post in posts)
    format_counts = {
        post_format.value: format_counts_counter.get(post_format.value, 0)
        for post_format in PostFormat
    }
    format_distribution = {
        name: round(count / len(posts), 6) if posts else 0.0
        for name, count in format_counts.items()
    }
    timestamps = sorted(
        (post.timestamp for post in posts if post.timestamp is not None), reverse=True
    )
    latest = timestamps[0] if timestamps else None
    recency = max(0.0, (as_of - latest).total_seconds() / 86_400) if latest else None
    intervals = [
        (timestamps[index] - timestamps[index + 1]).total_seconds() / 86_400
        for index in range(len(timestamps) - 1)
    ]
    span_days = (
        (timestamps[0] - timestamps[-1]).total_seconds() / 86_400
        if len(timestamps) >= 2
        else None
    )
    posts_per_week = (
        ((len(timestamps) - 1) / span_days) * 7
        if span_days is not None and span_days > 0
        else None
    )
    return ProfileMetrics(
        followers=profile.followers,
        posts_count=profile.posts_count,
        sampled_posts=len(posts),
        median_likes=safe_median(likes),
        median_comments=safe_median(comments),
        engagement_rate_pct=safe_median(engagement_rates),
        usable_engagement_posts=len(engagement_rates),
        engagement_confidence=(
            len(engagement_rates) / len(posts) if posts else None
        ),
        format_counts=format_counts,
        format_distribution=format_distribution,
        short_video_share=(
            format_counts[PostFormat.SHORT_VIDEO.value] / len(posts) if posts else None
        ),
        latest_post_at=latest,
        posting_recency_days=recency,
        posts_per_week=posts_per_week,
        median_post_interval_days=safe_median(intervals),
        signals=detect_signals(profile),
    )


def round_score(value: float) -> float:
    return round(min(max(value, 0.0), 100.0), 2)


def average_tie_percentile(values: dict[int, float]) -> dict[int, float]:
    if not values:
        return {}
    ordered = sorted(values.items(), key=lambda item: item[1])
    if len(ordered) == 1:
        return {ordered[0][0]: 1.0}
    result: dict[int, float] = {}
    position = 0
    while position < len(ordered):
        end = position
        while end + 1 < len(ordered) and math.isclose(
            ordered[end + 1][1], ordered[position][1], rel_tol=1e-12, abs_tol=1e-12
        ):
            end += 1
        average_index = (position + end) / 2
        percentile = average_index / (len(ordered) - 1)
        for index in range(position, end + 1):
            result[ordered[index][0]] = percentile
        position = end + 1
    return result


def audience_barter_points(followers: int | None) -> tuple[float, str]:
    if followers is None:
        return 0.0, "followers unavailable"
    if followers <= 1_000:
        return 4.0, "nano audience (≤1k)"
    if followers <= 10_000:
        return 5.0, "small creator audience (1k–10k)"
    if followers <= 50_000:
        return 4.0, "micro creator audience (10k–50k)"
    if followers <= 150_000:
        return 3.0, "mid-size audience (50k–150k)"
    if followers <= 300_000:
        return 2.0, "large audience (150k–300k)"
    return 1.0, "very large audience (>300k)"


def recency_points(recency_days: float | None) -> float:
    if recency_days is None:
        return 0.0
    if recency_days <= 7:
        return 6.0
    if recency_days <= 14:
        return 5.0
    if recency_days <= 30:
        return 4.0
    if recency_days <= 60:
        return 2.0
    if recency_days <= 90:
        return 1.0
    return 0.0


def frequency_points(posts_per_week: float | None) -> float:
    return min(4.0, max(0.0, ((posts_per_week or 0.0) / 2.0) * 4.0))


def score_profiles(
    profiles: list[Profile],
    metrics: dict[int, ProfileMetrics],
    brand_short_video_reference: float | None,
) -> dict[int, ProfileScore]:
    eligible = [
        profile for profile in profiles if profile.status == ProfileStatus.CREATOR
    ]
    engagement_values = {
        profile.source_index: metrics[profile.source_index].engagement_rate_pct
        for profile in eligible
        if metrics[profile.source_index].engagement_rate_pct is not None
    }
    engagement_percentiles = average_tie_percentile(engagement_values)  # type: ignore[arg-type]
    scores: dict[int, ProfileScore] = {}
    for profile in eligible:
        profile_metrics = metrics[profile.source_index]
        signals = profile_metrics.signals

        topic_weights = {
            "fashion": 7.0,
            "beauty": 5.0,
            "lifestyle": 4.0,
            "ugc": 5.0,
            "marketplace": 3.0,
        }
        topic_points = sum(
            weight for name, weight in topic_weights.items() if signals[name].detected
        )
        format_points = 0.0
        if (profile_metrics.short_video_share or 0) > 0:
            format_points += 3.0
        non_short_formats = sum(
            count
            for name, count in profile_metrics.format_counts.items()
            if name not in {PostFormat.SHORT_VIDEO.value, PostFormat.UNKNOWN.value}
        )
        if non_short_formats > 0:
            format_points += 1.0
        if any(post.caption.strip() for post in profile.latest_posts):
            format_points += 1.0
        brand_alignment = 0.0
        if (
            brand_short_video_reference is not None
            and profile_metrics.short_video_share is not None
        ):
            brand_alignment = 1.0 * (
                1
                - abs(
                    profile_metrics.short_video_share - brand_short_video_reference
                )
            )
        content_score = min(30.0, topic_points + format_points + brand_alignment)
        detected_topics = [
            name for name in topic_weights if signals[name].detected
        ]
        content_explanation = (
            f"{topic_points:.2f}/24 topical points from "
            f"{', '.join(detected_topics) if detected_topics else 'no matched target themes'}; "
            f"{format_points:.2f}/5 format-readiness points; "
            f"{brand_alignment:.2f}/1 format-alignment point against Nike/Apple."
        )

        integration_posts = int(
            signals["native_product_integration"].matches[0]
            if signals["native_product_integration"].matches
            else 0
        )
        integration_ratio = (
            integration_posts / profile_metrics.sampled_posts
            if profile_metrics.sampled_posts
            else 0.0
        )
        integration_score = (
            (6.0 if signals["ugc"].detected else 0.0)
            + (4.0 if signals["marketplace"].detected else 0.0)
            + (4.0 if signals["commercial_pr"].detected else 0.0)
            + min(4.0, integration_ratio * 8.0)
            + (2.0 if signals["contact"].detected else 0.0)
        )
        integration_score = min(20.0, integration_score)
        integration_explanation = (
            f"UGC={signals['ugc'].detected}, marketplace={signals['marketplace'].detected}, "
            f"commercial/PR={signals['commercial_pr'].detected}, "
            f"contact={signals['contact'].detected}; "
            f"{integration_posts}/{profile_metrics.sampled_posts} sampled posts show native "
            "integration evidence."
        )

        short_share = profile_metrics.short_video_share or 0.0
        short_score = min(15.0, short_share * 15.0)
        short_explanation = (
            f"Short video is {short_share:.1%} of {profile_metrics.sampled_posts} "
            "sampled posts; score is share × 15."
        )

        engagement_percentile = engagement_percentiles.get(profile.source_index)
        raw_engagement_score = (
            engagement_percentile * 15.0
            if engagement_percentile is not None
            else 0.0
        )
        engagement_confidence = profile_metrics.engagement_confidence or 0.0
        engagement_score = raw_engagement_score * engagement_confidence
        engagement_explanation = (
            f"Median per-post engagement rate is "
            f"{profile_metrics.engagement_rate_pct:.3f}% and ranks at the "
            f"{engagement_percentile:.1%} percentile of eligible creators. "
            f"Raw engagement component {raw_engagement_score:.2f}/15 × "
            f"confidence {profile_metrics.usable_engagement_posts}/"
            f"{profile_metrics.sampled_posts} ({engagement_confidence:.3f}) = "
            f"{engagement_score:.2f}/15 adjusted."
            if engagement_percentile is not None
            and profile_metrics.engagement_rate_pct is not None
            else "Engagement could not be calculated; no points assigned."
        )

        audience_points, audience_label = audience_barter_points(profile.followers)
        barter_score = (
            audience_points
            + (2.0 if signals["commercial_pr"].detected else 0.0)
            + (2.0 if signals["ugc"].detected else 0.0)
            + (1.0 if signals["marketplace"].detected else 0.0)
        )
        barter_score = min(10.0, barter_score)
        barter_explanation = (
            f"{audience_points:.1f}/5 for {audience_label}; +2 commercial/PR if present, "
            "+2 UGC if present, +1 marketplace if present."
        )

        recency = profile_metrics.posting_recency_days
        recency_score = recency_points(recency)
        frequency_score = frequency_points(profile_metrics.posts_per_week)
        activity_score = recency_score + frequency_score
        activity_explanation = (
            f"{recency_score:.1f}/6 from recency "
            f"({recency:.1f} days)" if recency is not None else "0/6: recency unavailable"
        )
        activity_explanation += (
            f"; {frequency_score:.1f}/4 from "
            f"{profile_metrics.posts_per_week:.2f} sampled posts/week."
            if profile_metrics.posts_per_week is not None
            else "; 0/4: frequency unavailable."
        )

        components = (
            ScoreComponent(
                "content_and_aesthetic_fit",
                round(content_score, 2),
                30.0,
                content_explanation,
            ),
            ScoreComponent(
                "native_product_integration_potential",
                round(integration_score, 2),
                20.0,
                integration_explanation,
            ),
            ScoreComponent(
                "short_video_consistency",
                round(short_score, 2),
                15.0,
                short_explanation,
            ),
            ScoreComponent(
                "engagement",
                round(engagement_score, 2),
                15.0,
                engagement_explanation,
            ),
            ScoreComponent(
                "barter_feasibility",
                round(barter_score, 2),
                10.0,
                barter_explanation,
            ),
            ScoreComponent(
                "recent_activity",
                round(activity_score, 2),
                10.0,
                activity_explanation,
            ),
        )
        total = round(sum(component.score for component in components), 2)
        strongest = max(components, key=lambda component: component.score / component.maximum)
        weakest = min(components, key=lambda component: component.score / component.maximum)
        explanation = (
            f"Total {total:.2f}/100. Strongest relative component: "
            f"{strongest.name} ({strongest.score:.2f}/{strongest.maximum:.0f}); "
            f"lowest relative component: {weakest.name} "
            f"({weakest.score:.2f}/{weakest.maximum:.0f})."
        )
        scores[profile.source_index] = ProfileScore(
            username=profile.username or f"record_{profile.source_index}",
            total=total,
            components=components,
            explanation=explanation,
            raw_engagement_score=round(raw_engagement_score, 2),
            engagement_confidence=round(engagement_confidence, 6),
        )
    return scores


def percentile(values: list[float], fraction: float) -> float | None:
    if not values:
        return None
    ordered = sorted(values)
    position = (len(ordered) - 1) * fraction
    lower = math.floor(position)
    upper = math.ceil(position)
    if lower == upper:
        return ordered[lower]
    return ordered[lower] + (ordered[upper] - ordered[lower]) * (position - lower)


def distribution_summary(values: Iterable[float | int | None]) -> dict[str, float | None]:
    cleaned = [float(value) for value in values if value is not None]
    return {
        "q1": round(percentile(cleaned, 0.25), 4) if cleaned else None,
        "median": round(percentile(cleaned, 0.5), 4) if cleaned else None,
        "q3": round(percentile(cleaned, 0.75), 4) if cleaned else None,
    }


def build_ideal_creator_profile(
    profiles: list[Profile],
    metrics: dict[int, ProfileMetrics],
    scores: dict[int, ProfileScore],
) -> dict[str, Any]:
    creators = [
        profile for profile in profiles if profile.status == ProfileStatus.CREATOR
    ]
    creator_metrics = [metrics[profile.source_index] for profile in creators]
    signal_names = [
        "commercial_pr",
        "contact",
        "fashion",
        "beauty",
        "lifestyle",
        "ugc",
        "marketplace",
        "native_product_integration",
    ]
    prevalence = {
        name: round(
            sum(item.signals[name].detected for item in creator_metrics)
            / len(creator_metrics),
            4,
        )
        if creator_metrics
        else None
        for name in signal_names
    }
    score_values = [score.total for score in scores.values()]
    return {
        "basis": (
            "Robust cohort medians and interquartile ranges from quantitatively eligible "
            "creators only. Nike and Apple influence only the documented format-alignment "
            "signal; they do not enter creator distributions."
        ),
        "eligible_creator_count": len(creators),
        "audience": {
            "followers": distribution_summary(item.followers for item in creator_metrics),
            "recommended_interpretation": (
                "Prioritize creators near the cohort median or within the interquartile "
                "range; smaller audiences improve barter feasibility while larger audiences "
                "may increase reach and cost."
            ),
        },
        "performance": {
            "median_likes": distribution_summary(
                item.median_likes for item in creator_metrics
            ),
            "median_comments": distribution_summary(
                item.median_comments for item in creator_metrics
            ),
            "engagement_rate_pct": distribution_summary(
                item.engagement_rate_pct for item in creator_metrics
            ),
            "engagement_confidence": distribution_summary(
                item.engagement_confidence for item in creator_metrics
            ),
        },
        "activity": {
            "posting_recency_days": distribution_summary(
                item.posting_recency_days for item in creator_metrics
            ),
            "posts_per_week": distribution_summary(
                item.posts_per_week for item in creator_metrics
            ),
        },
        "formats": {
            "short_video_share": distribution_summary(
                item.short_video_share for item in creator_metrics
            ),
            "ideal": (
                "Consistent short-form video with enough caption, mention, tag, or "
                "marketplace evidence to demonstrate native product integration."
            ),
        },
        "signal_prevalence": prevalence,
        "recommended_signals": [
            name for name, share in prevalence.items() if share is not None and share >= 0.35
        ],
        "barter_feasibility": {
            "preferred_evidence": [
                "UGC positioning",
                "commercial or PR wording",
                "marketplace/product-review experience",
                "small-to-mid-size audience",
                "recent posting activity",
            ],
            "score_distribution": distribution_summary(score_values),
        },
        "selection_principle": (
            "Use the score as a transparent shortlist aid, then visually review content "
            "quality, brand safety, audience geography, authenticity, and commercial terms."
        ),
    }


def reconcile_workbook(
    references: list[WorkbookReference],
    profiles: list[Profile],
    manual_audit: dict[str, Any],
) -> dict[str, Any]:
    recovery = manual_audit.get("manual_human_in_the_loop_recovery", {})
    corrections = {
        str(item.get("source_username", "")).casefold(): str(
            item.get("verified_username", "")
        ).casefold()
        for item in recovery.get("confirmed_corrections", [])
        if isinstance(item, dict)
        and item.get("source_username")
        and item.get("verified_username")
    }
    workbook_names = [
        reference.selected_username
        for reference in references
        if reference.selected_username
    ]
    corrected_workbook_names = [corrections.get(name, name) for name in workbook_names]
    json_names = {
        (
            profile.username
            or normalize_username_from_url(profile.input_url)
            or ""
        ).casefold()
        for profile in profiles
    }
    unresolved = {
        str(value).casefold() for value in manual_audit.get("remaining_unresolved", [])
    }
    unmatched_after_audit = sorted(
        name
        for name in corrected_workbook_names
        if name not in json_names and name not in unresolved
    )
    return {
        "workbook_reference_count": len(references),
        "json_record_count": len(profiles),
        "hyperlink_target_used_count": sum(
            reference.hyperlink_preferred for reference in references
        ),
        "display_target_mismatch_count": sum(
            reference.display_target_mismatch for reference in references
        ),
        "display_target_mismatches": [
            reference.serializable()
            for reference in references
            if reference.display_target_mismatch
        ],
        "manual_correction_count": len(corrections),
        "manual_corrections": corrections,
        "unmatched_after_manual_audit": unmatched_after_audit,
        "reconciled": len(references) == len(profiles) and not unmatched_after_audit,
        "rule": (
            "Use the Excel cell hyperlink target when present; the displayed cell text "
            "can be truncated or stale. Then apply only explicitly verified manual corrections."
        ),
    }


def determine_as_of(profiles: list[Profile], requested: datetime | None) -> datetime:
    if requested is not None:
        return requested.astimezone(timezone.utc)
    env_value = os.getenv("ANALYSIS_AS_OF", "").strip()
    if env_value:
        parsed = datetime.fromisoformat(env_value.replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone.utc)
        return parsed.astimezone(timezone.utc)
    timestamps = [
        post.timestamp
        for profile in profiles
        for post in profile.latest_posts
        if post.timestamp is not None
    ]
    if timestamps:
        return max(timestamps)
    return datetime.now(timezone.utc)


def analyze_dataset(
    profiles_path: Path,
    audit_path: Path,
    workbook_path: Path,
    *,
    as_of: datetime | None = None,
) -> AnalysisResult:
    raw_profiles = read_json(profiles_path)
    if not isinstance(raw_profiles, list):
        raise ValueError(f"{profiles_path}: top-level JSON value must be an array")
    raw_audit = read_json(audit_path)
    if not isinstance(raw_audit, dict):
        raise ValueError(f"{audit_path}: top-level JSON value must be an object")
    profiles = [Profile.from_raw(raw, index) for index, raw in enumerate(raw_profiles)]
    for profile in profiles:
        classify_profile(profile)
    analysis_as_of = determine_as_of(profiles, as_of)
    metrics = {
        profile.source_index: calculate_metrics(profile, analysis_as_of)
        for profile in profiles
    }
    brand_profiles = [
        profile
        for profile in profiles
        if profile.status == ProfileStatus.BRAND_REFERENCE
    ]
    brand_short_video_reference = safe_median(
        metrics[profile.source_index].short_video_share for profile in brand_profiles
    )
    scores = score_profiles(
        profiles, metrics, brand_short_video_reference
    )
    workbook_references = read_workbook_references(workbook_path)
    workbook_reconciliation = reconcile_workbook(
        workbook_references, profiles, raw_audit
    )
    counts = Counter(
        profile.status.value if profile.status else "unclassified" for profile in profiles
    )
    validation_issues = [
        issue for profile in profiles for issue in profile.validation_issues
    ]
    usernames = [
        profile.username.casefold() for profile in profiles if profile.username
    ]
    duplicates = sorted(
        username for username, count in Counter(usernames).items() if count > 1
    )
    dataset_summary = {
        "total_records": len(profiles),
        "counts_by_status": dict(sorted(counts.items())),
        "quantitatively_analyzed_creators": counts.get(ProfileStatus.CREATOR.value, 0),
        "scored_creators": len(scores),
        "validation_issue_count": len(validation_issues),
        "duplicate_usernames": duplicates,
        "analysis_as_of": analysis_as_of.isoformat(),
    }
    brand_reference_signals = [
        {
            "username": profile.username,
            "role": "visual_and_brand_reference_only",
            "sampled_posts": metrics[profile.source_index].sampled_posts,
            "format_distribution": metrics[
                profile.source_index
            ].format_distribution,
            "short_video_share": metrics[profile.source_index].short_video_share,
            "signals": {
                name: signal.serializable()
                for name, signal in metrics[profile.source_index].signals.items()
            },
            "excluded_from_creator_statistics": True,
        }
        for profile in brand_profiles
    ]
    ideal = build_ideal_creator_profile(profiles, metrics, scores)
    return AnalysisResult(
        generated_at=datetime.now(timezone.utc),
        analysis_as_of=analysis_as_of,
        profiles=profiles,
        metrics=metrics,
        scores=scores,
        workbook_references=workbook_references,
        workbook_reconciliation=workbook_reconciliation,
        manual_audit=raw_audit,
        dataset_summary=dataset_summary,
        brand_reference_signals=brand_reference_signals,
        ideal_creator_profile=ideal,
        validation_issues=validation_issues,
    )
