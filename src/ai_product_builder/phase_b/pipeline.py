"""End-to-end orchestration for the Phase B Instagram MVP."""

from __future__ import annotations

import csv
import hashlib
import json
from collections import Counter
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, replace
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .account_types import (
    AccountTypeAssessment,
    assess_account_type,
    with_account_theme_evidence,
)
from .barter_signals import assess_barter_signals
from .compatibility import assess_compatibility
from .config import (
    PhaseBConfig,
    load_env_secrets,
    load_phase_b_config,
    require_secret,
)
from .eligibility import (
    evaluate_candidate_eligibility,
    is_instagram_post_url,
)
from .enrichment import (
    calculate_candidate_metrics,
    known_format_posts,
    normalized_post_format,
    resolve_as_of,
)
from .errors import (
    ConfigurationError,
    EvidenceValidationError,
    InputValidationError,
    InsufficientCandidatePoolError,
    OutputWriteError,
    PhaseBError,
)
from .evidence import collect_signal_evidence, flatten_evidence
from .exclusions import ExclusionMatch, ExclusionRegistry
from .io.artifacts import (
    ARTIFACT_FILENAMES,
    generate_phase_b_failure_artifacts,
    generate_phase_b_artifacts,
    run_directory,
)
from .io.google_sheets import (
    GoogleSheetsAdapter,
    GoogleSheetsWriteReceipt,
    build_google_sheets_client,
)
from .models import (
    CandidateIdentity,
    CandidateResult,
    CreatorProfile,
    DiscoveryHit,
    EligibilityDecision,
    PhaseBRunManifest,
    QuerySpec,
    SignalEvidence,
    parse_datetime,
)
from .near_miss import build_near_miss_candidates
from .normalization import canonicalize_profile_url, normalize_username
from .offers import generate_deterministic_offer
from .providers.apify import ApifyInstagramProvider
from .providers.base import InstagramProvider
from .providers.fixtures import FixtureInstagramProvider
from .queries import generate_queries
from .scoring import (
    calculate_discovery_confidence,
    load_phase_a_reference,
    score_candidate,
)
from .selection import (
    rank_candidates,
    select_top_candidates,
    selection_explanation,
)


@dataclass(frozen=True, slots=True)
class DeduplicationOutcome:
    identities: tuple[CandidateIdentity, ...]
    report: Mapping[str, Any]
    excluded_records: tuple[Mapping[str, Any], ...]


@dataclass(frozen=True, slots=True)
class PhaseBRunResult:
    manifest: PhaseBRunManifest
    run_dir: Path
    generated_paths: tuple[Path, ...]
    selected_candidates: tuple[CandidateResult, ...]
    eligible_candidates: tuple[CandidateResult, ...]
    google_sheets_receipt: GoogleSheetsWriteReceipt | None = None


@dataclass(frozen=True, slots=True)
class LiveCampaignAssessment:
    target_theme_relevant: bool
    language_compatible: bool
    geography_compatible: bool
    delivery_market_review_required: bool
    commercial_conflict: bool
    own_fashion_brand: bool
    own_clothing_store_or_showroom: bool
    explicit_no_barter: bool
    blogger_content_sufficient: bool
    short_video_ready: bool
    recent_fashion_posts: int
    recent_short_video_posts: int
    recent_fashion_post_urls: tuple[str, ...]
    manual_review_required: bool
    preferred_audience_range: bool
    audience_review_required: bool
    explanation: str

    def to_dict(self) -> dict[str, Any]:
        return {
            field: getattr(self, field)
            for field in self.__dataclass_fields__
        }


def _live_campaign_assessment(
    profile: CreatorProfile,
    evidence: Mapping[str, Iterable[SignalEvidence]],
    account: AccountTypeAssessment,
    *,
    language_compatible: bool,
    detected_geography: str | None,
    delivery_market_review_required: bool,
    delivery_market_conflict: bool,
    explicit_no_barter: bool,
    as_of: datetime | None,
    maximum_recency_days: float,
    minimum_recent_fashion_posts: int,
    minimum_short_video_posts: int,
    preferred_followers_min: int,
    preferred_followers_max: int,
) -> LiveCampaignAssessment:
    effective_as_of = resolve_as_of(profile, as_of)
    recent_posts: dict[str, Any] = {}
    recent_short_video_posts = 0
    for index, post in enumerate(profile.recent_posts):
        if post.timestamp is None or not is_instagram_post_url(post.url):
            continue
        timestamp = (
            post.timestamp.replace(tzinfo=effective_as_of.tzinfo)
            if post.timestamp.tzinfo is None
            else post.timestamp.astimezone(effective_as_of.tzinfo)
        )
        age_days = (
            effective_as_of - timestamp
        ).total_seconds() / 86_400
        if not 0 <= age_days <= maximum_recency_days:
            continue
        reference = post.post_id or post.url or f"index:{index}"
        recent_posts[reference] = post
        if normalized_post_format(post) == "short_video":
            recent_short_video_posts += 1

    fashion_references = {
        item.source_reference
        for item in evidence.get("fashion", ())
        if item.observation_type == "direct"
        and item.source_field.startswith("recent_posts.")
        and item.source_reference in recent_posts
    }
    recent_fashion_urls = tuple(
        dict.fromkeys(
            str(recent_posts[reference].url)
            for reference in sorted(
                fashion_references,
                key=lambda item: recent_posts[item].timestamp,
                reverse=True,
            )
            if recent_posts[reference].url
        )
    )
    geography_compatible = bool(
        detected_geography
        and "россия" in detected_geography.casefold()
        and not delivery_market_conflict
        and not delivery_market_review_required
    )
    short_video_ready = (
        recent_short_video_posts >= minimum_short_video_posts
    )
    direct_fashion_sufficient = (
        len(recent_fashion_urls) >= minimum_recent_fashion_posts
    )
    blogger_content_sufficient = bool(
        account.account_type.value == "personal_creator"
        and not account.professional_portfolio
        and direct_fashion_sufficient
        and short_video_ready
    )
    followers = profile.followers
    preferred_audience_range = bool(
        followers is not None
        and preferred_followers_min <= followers <= preferred_followers_max
    )
    audience_review_required = bool(
        followers is not None and followers > preferred_followers_max
    )
    explanation = (
        f"Direct recent fashion posts: {len(recent_fashion_urls)}/"
        f"{minimum_recent_fashion_posts}; recent short-video posts: "
        f"{recent_short_video_posts}/{minimum_short_video_posts}; "
        f"language compatible: {language_compatible}; directly evidenced "
        f"Russia compatibility: {geography_compatible}; account type: "
        f"{account.account_type.value}; professional portfolio: "
        f"{account.professional_portfolio}; commercial conflict: "
        f"{account.commercial_conflict}."
    )
    return LiveCampaignAssessment(
        target_theme_relevant=account.theme_relevant,
        language_compatible=language_compatible,
        geography_compatible=geography_compatible,
        delivery_market_review_required=(
            delivery_market_review_required
        ),
        commercial_conflict=account.commercial_conflict,
        own_fashion_brand=account.own_fashion_brand,
        own_clothing_store_or_showroom=(
            account.own_clothing_store_or_showroom
        ),
        explicit_no_barter=explicit_no_barter,
        blogger_content_sufficient=blogger_content_sufficient,
        short_video_ready=short_video_ready,
        recent_fashion_posts=len(recent_fashion_urls),
        recent_short_video_posts=recent_short_video_posts,
        recent_fashion_post_urls=recent_fashion_urls,
        manual_review_required=True,
        preferred_audience_range=preferred_audience_range,
        audience_review_required=audience_review_required,
        explanation=explanation,
    )


def _with_run_level_exclusions(
    registry: ExclusionRegistry,
    config: PhaseBConfig,
) -> ExclusionRegistry:
    matches = []
    for item in config.run_level_exclusions:
        normalized = normalize_username(item.username)
        canonical = canonicalize_profile_url(item.profile_url)
        if normalized is None or canonical is None:
            raise ConfigurationError(
                "run-level exclusion contains a malformed Instagram identity"
            )
        matches.append(
            ExclusionMatch(
                normalized_username=normalized,
                canonical_profile_url=canonical,
                reasons=(item.reason,),
                source_values=(item.username, item.profile_url),
            )
        )
    return registry.extended(matches)


def _phase_a_follower_outer_fence(
    ideal_profile_document: Mapping[str, Any],
) -> tuple[float, float, float]:
    try:
        followers = ideal_profile_document["ideal_creator_profile"][
            "audience"
        ]["followers"]
        q1 = float(followers["q1"])
        q3 = float(followers["q3"])
    except (KeyError, TypeError, ValueError) as exc:
        raise InputValidationError(
            "ideal creator profile is missing follower q1/q3"
        ) from exc
    if q1 < 0 or q3 < q1:
        raise InputValidationError(
            "ideal creator profile has invalid follower q1/q3"
        )
    return q1, q3, q3 + 3.0 * (q3 - q1)


def _apply_live_campaign_guards(
    decision: EligibilityDecision,
    *,
    explicit_barter_refusal: bool,
    language_compatible: bool,
    delivery_market_conflict: bool,
    high_audience: bool,
    geography_compatible: bool = True,
    commercial_conflict: bool = False,
    own_fashion_brand: bool = False,
    own_clothing_store_or_showroom: bool = False,
    blogger_content_sufficient: bool = True,
    short_video_ready: bool = True,
    recent_fashion_posts: int = 3,
    minimum_recent_fashion_posts: int = 3,
) -> EligibilityDecision:
    """Apply campaign-only blockers after the generic eligibility checks."""

    if not decision.eligible:
        return decision
    blockers: list[str] = []
    if explicit_barter_refusal:
        blockers.append("explicit_no_barter_statement")
    if not language_compatible:
        blockers.append("campaign_language_mismatch")
    if delivery_market_conflict:
        blockers.append("delivery_market_conflict")
    elif not geography_compatible:
        blockers.append("delivery_market_compatibility_unverified")
    if commercial_conflict:
        if own_fashion_brand:
            blockers.append("commercial_conflict_own_fashion_brand")
        if own_clothing_store_or_showroom:
            blockers.append(
                "commercial_conflict_clothing_store_or_showroom"
            )
        if not (own_fashion_brand or own_clothing_store_or_showroom):
            blockers.append("commercial_conflict")
    if not short_video_ready:
        blockers.append("short_video_evidence_missing")
    if recent_fashion_posts < minimum_recent_fashion_posts:
        blockers.append(
            "insufficient_recent_fashion_posts:"
            f"{recent_fashion_posts}<{minimum_recent_fashion_posts}"
        )
    if not blogger_content_sufficient:
        blockers.append("blogger_content_insufficient")
    if blockers:
        return EligibilityDecision(
            eligible=False,
            status="ineligible",
            reasons=tuple(dict.fromkeys(blockers)),
        )
    if high_audience:
        return EligibilityDecision(
            eligible=False,
            status="needs_review",
            reasons=("high_audience_barter_review_required",),
        )
    return decision


def validate_phase_b_config(
    config_path: Path,
    *,
    expected_mode: str | None = None,
) -> tuple[PhaseBConfig, tuple[str, ...]]:
    """Load configuration and validate all non-secret local dependencies."""

    config_path = Path(config_path)
    config = load_phase_b_config(config_path)
    if expected_mode is not None and config.mode != expected_mode:
        raise ConfigurationError(
            f"Command mode '{expected_mode}' does not match config mode "
            f"'{config.mode}'."
        )
    if not 3 <= config.discovery.final_count <= 5:
        raise ConfigurationError(
            "discovery.final_count must be between 3 and 5 for the MVP"
        )
    if config.discovery.minimum_final_count != 3:
        raise ConfigurationError(
            "discovery.minimum_final_count must be 3 for the MVP"
        )
    if config.discovery.minimum_unique_pool <= config.discovery.final_count:
        raise ConfigurationError(
            "discovery.minimum_unique_pool must be larger than final_count"
        )
    if config.eligibility.minimum_usable_posts < 6:
        raise ConfigurationError(
            "eligibility.minimum_usable_posts cannot be lower than 6"
        )
    if config.eligibility.maximum_recency_days > 90:
        raise ConfigurationError(
            "eligibility.maximum_recency_days cannot exceed 90"
        )
    if (
        config.eligibility.preferred_followers_min
        > config.eligibility.preferred_followers_max
    ):
        raise ConfigurationError(
            "eligibility.preferred_followers_min cannot exceed "
            "preferred_followers_max"
        )
    if len(config.discovery.query_texts) > 20:
        raise ConfigurationError(
            "discovery.query_texts cannot contain more than 20 queries"
        )
    if config.outputs.csv_encoding.casefold() != "utf-8-sig":
        raise ConfigurationError(
            "outputs.csv_encoding must be 'utf-8-sig' for the MVP"
        )
    if config.outputs.workbook_sheet != "Новые блоггеры":
        raise ConfigurationError(
            "outputs.workbook_sheet must be exactly 'Новые блоггеры'"
        )

    required_files = {
        "ideal_creator_profile": config.inputs.ideal_creator_profile,
        "source_analysis": config.inputs.source_analysis,
        "instagram_profiles": config.inputs.instagram_profiles,
        "manual_audit": config.inputs.manual_audit,
        "workbook": config.inputs.workbook,
    }
    if config.provider.type == "fixture":
        if config.provider.fixture is None:
            raise ConfigurationError("fixture provider configuration is missing")
        required_files.update(
            {
                "fixture.discovery_path": config.provider.fixture.discovery_path,
                "fixture.profiles_path": config.provider.fixture.profiles_path,
            }
        )
    missing = [
        f"{name}: {path}"
        for name, path in required_files.items()
        if not Path(path).is_file()
    ]
    if missing:
        raise InputValidationError(
            "Required Phase B input file(s) are missing.",
            details={"missing": missing},
        )

    ideal = _read_json_object(
        config.inputs.ideal_creator_profile, "ideal_creator_profile"
    )
    # Query construction is also a schema validation pass.
    generate_queries(
        ideal,
        config.campaign,
        query_texts=config.discovery.query_texts,
    )
    _validate_source_analysis(config.inputs.source_analysis)
    _with_run_level_exclusions(
        ExclusionRegistry.from_files(
            config.inputs.instagram_profiles,
            config.inputs.manual_audit,
        ),
        config,
    )
    _validate_workbook(config.inputs.workbook)

    if config.analysis_as_of is not None and parse_datetime(config.analysis_as_of) is None:
        raise ConfigurationError("analysis_as_of must be a valid ISO-8601 timestamp")

    warnings: list[str] = []
    if config.provider.type == "fixture":
        fixture = config.provider.fixture
        assert fixture is not None
        provider = FixtureInstagramProvider(
            fixture.discovery_path, fixture.profiles_path
        )
        fixture_hits = provider.discover(())
        if len(fixture_hits) < 30:
            raise InputValidationError(
                "Demo discovery fixture must contain at least 30 records.",
                details={"record_count": len(fixture_hits)},
            )
    else:
        apify = config.provider.apify
        if apify is None:
            raise ConfigurationError("Apify provider configuration is missing")
        if apify.maximum_items < config.discovery.target_pool_size:
            raise ConfigurationError(
                "provider.apify.maximum_items cannot be lower than "
                "discovery.target_pool_size"
            )
        if (
            2.0 * apify.max_total_charge_usd
            > apify.max_combined_charge_usd + 1e-9
        ):
            raise ConfigurationError(
                "The combined Apify budget cannot cover one discovery and "
                "one enrichment run at their configured server-side caps"
            )
        # Construction validates required mappings without making a request.
        ApifyInstagramProvider(apify, "configuration-validation-only")
        secrets = load_env_secrets(config.env_file)
        if not secrets.get("APIFY_TOKEN"):
            warnings.append(
                f"Live execution requires APIFY_TOKEN in {config.env_file}."
            )
        if (
            apify.discovery_actor_id.startswith("REPLACE_WITH_")
            or apify.enrichment_actor_id.startswith("REPLACE_WITH_")
        ):
            warnings.append(
                "Live actor IDs are placeholders and must be replaced before "
                "the first Apify smoke test."
            )

    if config.google_sheets.enabled:
        if not config.google_sheets.spreadsheet_id:
            raise ConfigurationError(
                "google_sheets.spreadsheet_id is required when enabled"
            )
        secrets = load_env_secrets(config.env_file)
        if not secrets.get("GOOGLE_SERVICE_ACCOUNT_JSON"):
            warnings.append(
                "Google Sheets is enabled, but GOOGLE_SERVICE_ACCOUNT_JSON is "
                f"missing from {config.env_file}."
            )
    return config, tuple(warnings)


def deduplicate_discovery_hits(
    hits: Sequence[DiscoveryHit],
    exclusion_registry: ExclusionRegistry,
    *,
    known_query_ids: Iterable[str] = (),
) -> DeduplicationOutcome:
    """Exclude frozen identities and merge discoveries by username/profile URL."""

    allowed_query_ids = frozenset(known_query_ids)
    prepared: list[dict[str, Any]] = []
    excluded: list[dict[str, Any]] = []
    parent: dict[str, str] = {}

    def find(value: str) -> str:
        parent.setdefault(value, value)
        while parent[value] != value:
            parent[value] = parent[parent[value]]
            value = parent[value]
        return value

    def union(left: str, right: str) -> None:
        left_root = find(left)
        right_root = find(right)
        if left_root != right_root:
            # Stable first-seen roots keep fixture runs byte-for-byte predictable.
            parent[right_root] = left_root

    for index, hit in enumerate(hits):
        username_key = normalize_username(hit.username)
        url_key = normalize_username(hit.profile_url)
        if hit.platform.casefold() != "instagram":
            excluded.append(
                _excluded_hit(hit, "unsupported_platform", hit.validation_issues)
            )
            continue
        if username_key is None and url_key is None:
            excluded.append(
                _excluded_hit(
                    hit,
                    "malformed_discovery_record",
                    hit.validation_issues
                    or ("missing or malformed Instagram identity",),
                )
            )
            continue
        exclusion = exclusion_registry.match(hit.username, hit.profile_url)
        if exclusion is not None:
            exclusion_reason = (
                "previous_live_run_exclusion"
                if "previous_live_run_exclusion" in exclusion.reasons
                else "source_exclusion"
            )
            excluded.append(
                _excluded_hit(
                    hit,
                    exclusion_reason,
                    exclusion.reasons,
                    exclusion_match=exclusion,
                )
            )
            continue

        tokens = tuple(dict.fromkeys(value for value in (username_key, url_key) if value))
        for token in tokens:
            find(token)
        if len(tokens) == 2:
            union(tokens[0], tokens[1])
        query_ids = tuple(
            dict.fromkeys(
                query_id
                for query_id in hit.query_ids
                if not allowed_query_ids or query_id in allowed_query_ids
            )
        )
        prepared.append(
            {
                "index": index,
                "hit": hit,
                "tokens": tokens,
                "query_ids": query_ids,
                "identity_conflict": len(tokens) > 1,
            }
        )

    grouped: dict[str, list[dict[str, Any]]] = {}
    for item in prepared:
        root = find(item["tokens"][0])
        grouped.setdefault(root, []).append(item)

    identities: list[CandidateIdentity] = []
    groups_report: list[dict[str, Any]] = []
    duplicate_count = 0
    for group in grouped.values():
        group.sort(key=lambda item: item["index"])
        first_hit: DiscoveryHit = group[0]["hit"]
        all_tokens = tuple(
            dict.fromkeys(token for item in group for token in item["tokens"])
        )
        normalized = (
            normalize_username(first_hit.username)
            or normalize_username(first_hit.profile_url)
            or all_tokens[0]
        )
        canonical = canonicalize_profile_url(normalized)
        if canonical is None:  # Defensive: prepared records always normalized.
            excluded.append(
                _excluded_hit(
                    first_hit,
                    "malformed_discovery_record",
                    ("could not construct a canonical profile URL",),
                )
            )
            continue
        username = (
            first_hit.username.strip()
            if normalize_username(first_hit.username) == normalized
            else normalized
        )
        query_ids = tuple(
            dict.fromkeys(
                query_id for item in group for query_id in item["query_ids"]
            )
        )
        provider_ids = tuple(
            dict.fromkeys(
                value
                for item in group
                for value in (
                    item["hit"].provider_id,
                    item["hit"].provider_run_id,
                )
                if value
            )
        )
        identity_conflict = (
            len(all_tokens) > 1
            or any(item["identity_conflict"] for item in group)
        )
        identity = CandidateIdentity(
            platform="instagram",
            username=username,
            normalized_username=normalized,
            profile_url=canonical,
            canonical_profile_url=canonical,
            query_ids=query_ids,
            provider_ids=provider_ids,
            identity_conflict=identity_conflict,
        )
        identities.append(identity)
        duplicate_count += max(0, len(group) - 1)
        for duplicate in group[1:]:
            excluded.append(
                _excluded_hit(
                    duplicate["hit"],
                    "duplicate_merged",
                    (f"merged_into:{normalized}",),
                )
            )
        groups_report.append(
            {
                "normalized_username": normalized,
                "canonical_profile_url": canonical,
                "merged_hit_count": len(group),
                "query_ids": list(query_ids),
                "provider_ids": list(provider_ids),
                "identity_conflict": identity_conflict,
            }
        )

    reason_counts = Counter(str(item["exclusion_reason"]) for item in excluded)
    report = {
        "raw_hit_count": len(hits),
        "processable_hit_count": len(prepared),
        "unique_candidate_count": len(identities),
        "duplicate_hit_count": duplicate_count,
        "excluded_record_count": len(excluded),
        "exclusion_reason_counts": dict(sorted(reason_counts.items())),
        "groups": groups_report,
    }
    return DeduplicationOutcome(
        identities=tuple(identities),
        report=report,
        excluded_records=tuple(excluded),
    )


def run_phase_b(
    config_path: Path,
    output_dir: Path,
    *,
    expected_mode: str | None = None,
    sheets_client: Any | None = None,
) -> PhaseBRunResult:
    """Run Phase B without any outreach or message-sending side effects."""

    config_path = Path(config_path)
    config, validation_warnings = validate_phase_b_config(
        config_path, expected_mode=expected_mode
    )
    config_digest = hashlib.sha256(config_path.read_bytes()).hexdigest()
    run_id = (
        "demo"
        if config.mode == "demo"
        else datetime.now(timezone.utc).strftime("live-%Y%m%dT%H%M%SZ")
    )
    manifest = PhaseBRunManifest.started(
        run_id=run_id,
        mode=config.mode,
        provider=config.provider.type,
        config_digest=config_digest,
    )
    manifest.warnings.extend(validation_warnings)
    run_dir = run_directory(Path(output_dir), manifest)
    run_dir.mkdir(parents=True, exist_ok=True)
    provider: InstagramProvider | None = None
    resolved_sheets_client = sheets_client
    if config.google_sheets.enabled and resolved_sheets_client is None:
        google_secret = load_env_secrets(config.env_file).get(
            "GOOGLE_SERVICE_ACCOUNT_JSON"
        )
        if google_secret:
            try:
                resolved_sheets_client = build_google_sheets_client(
                    google_secret
                )
            except (OSError, RuntimeError, ValueError) as exc:
                wrapped = ConfigurationError(
                    f"Could not configure Google Sheets: {exc}"
                )
                _record_failed_manifest(run_dir, manifest, wrapped)
                raise wrapped from exc

    try:
        ideal = _read_json_object(
            config.inputs.ideal_creator_profile, "ideal_creator_profile"
        )
        queries = generate_queries(
            ideal,
            config.campaign,
            query_texts=config.discovery.query_texts,
        )
        exclusion_registry = _with_run_level_exclusions(
            ExclusionRegistry.from_files(
                config.inputs.instagram_profiles,
                config.inputs.manual_audit,
            ),
            config,
        )
        phase_a_q1, phase_a_q3, barter_review_threshold = (
            _phase_a_follower_outer_fence(ideal)
        )
        provider = _build_provider(config)
        discovery_hits = provider.discover(queries)
        deduplicated = deduplicate_discovery_hits(
            discovery_hits,
            exclusion_registry,
            known_query_ids=(query.query_id for query in queries),
        )
        if (
            len(deduplicated.identities)
            < config.discovery.minimum_unique_pool
        ):
            message = (
                "Unique discovery pool is below the configured minimum: "
                f"{len(deduplicated.identities)} < "
                f"{config.discovery.minimum_unique_pool}."
            )
            if config.mode == "live":
                raise InsufficientCandidatePoolError(
                    message,
                    details={
                        "stage": "before_enrichment",
                        "unique_candidates": len(
                            deduplicated.identities
                        ),
                        "minimum_unique_pool": (
                            config.discovery.minimum_unique_pool
                        ),
                    },
                )
            manifest.warnings.append(message)

        profiles = provider.enrich(deduplicated.identities)
        as_of = (
            parse_datetime(config.analysis_as_of)
            if config.analysis_as_of is not None
            else None
        )
        reference = load_phase_a_reference(
            config.inputs.source_analysis,
            config.inputs.ideal_creator_profile,
        )

        excluded_records = list(deduplicated.excluded_records)
        enriched_records: list[dict[str, Any]] = []
        eligible_results: list[CandidateResult] = []
        provider_run_ids = {
            value
            for hit in discovery_hits
            for value in (hit.provider_run_id,)
            if value
        }
        for profile in profiles:
            provider_run_ids.update(profile.provider_run_ids)
            metrics = calculate_candidate_metrics(profile, as_of)
            account_assessment = assess_account_type(profile)
            evidence = with_account_theme_evidence(
                collect_signal_evidence(profile),
                account_assessment,
            )
            barter_assessment = assess_barter_signals(profile)
            compatibility = assess_compatibility(
                profile,
                config.campaign,
            )
            live_assessment = (
                _live_campaign_assessment(
                    profile,
                    evidence,
                    account_assessment,
                    language_compatible=(
                        compatibility.campaign_language_compatible
                    ),
                    detected_geography=(
                        compatibility.detected_geography
                    ),
                    delivery_market_review_required=(
                        compatibility.delivery_market_review_required
                    ),
                    delivery_market_conflict=(
                        compatibility.delivery_market_conflict
                    ),
                    explicit_no_barter=(
                        barter_assessment.explicit_refusal
                    ),
                    as_of=as_of,
                    maximum_recency_days=(
                        config.eligibility.maximum_recency_days
                    ),
                    minimum_recent_fashion_posts=(
                        config.eligibility.minimum_recent_fashion_posts
                    ),
                    minimum_short_video_posts=(
                        config.eligibility.minimum_short_video_posts
                    ),
                    preferred_followers_min=(
                        config.eligibility.preferred_followers_min
                    ),
                    preferred_followers_max=(
                        config.eligibility.preferred_followers_max
                    ),
                )
                if config.mode == "live"
                else None
            )
            decision = evaluate_candidate_eligibility(
                profile,
                metrics,
                evidence,
                as_of=as_of,
                max_recency_days=config.eligibility.maximum_recency_days,
                minimum_usable_posts=config.eligibility.minimum_usable_posts,
                account_assessment=account_assessment,
            )
            high_audience = bool(
                metrics.followers is not None
                and metrics.followers > barter_review_threshold
            )
            if config.mode == "live":
                decision = _apply_live_campaign_guards(
                    decision,
                    explicit_barter_refusal=(
                        barter_assessment.explicit_refusal
                    ),
                    language_compatible=(
                        compatibility.campaign_language_compatible
                    ),
                    delivery_market_conflict=(
                        compatibility.delivery_market_conflict
                    ),
                    high_audience=high_audience,
                    geography_compatible=(
                        live_assessment.geography_compatible
                    ),
                    commercial_conflict=(
                        live_assessment.commercial_conflict
                    ),
                    own_fashion_brand=(
                        live_assessment.own_fashion_brand
                    ),
                    own_clothing_store_or_showroom=(
                        live_assessment.own_clothing_store_or_showroom
                    ),
                    blogger_content_sufficient=(
                        live_assessment.blogger_content_sufficient
                    ),
                    short_video_ready=(
                        live_assessment.short_video_ready
                    ),
                    recent_fashion_posts=(
                        live_assessment.recent_fashion_posts
                    ),
                    minimum_recent_fashion_posts=(
                        config.eligibility.minimum_recent_fashion_posts
                    ),
                )
            discovery_confidence = calculate_discovery_confidence(
                profile, evidence
            )
            score_preview = (
                score_candidate(profile, metrics, evidence, reference)
                if metrics.usable_posts >= 6 and metrics.sampled_posts > 0
                else None
            )
            offer = None
            if decision.eligible:
                try:
                    offer = generate_deterministic_offer(
                        profile,
                        config.campaign,
                        evidence,
                        as_of=as_of,
                        max_age_days=config.eligibility.maximum_recency_days,
                        required_signal_types=(
                            ("fashion",)
                            if config.mode == "live"
                            else None
                        ),
                    )
                except EvidenceValidationError as exc:
                    decision = EligibilityDecision(
                        eligible=False,
                        status="ineligible",
                        reasons=(
                            "offer_grounding_evidence_missing",
                            str(exc),
                        ),
                    )

            enriched_records.append(
                {
                    "profile": profile.to_dict(),
                    "metrics": metrics.to_dict(),
                    "evidence": {
                        name: [item.to_dict() for item in items]
                        for name, items in evidence.items()
                    },
                    "account_type_assessment": account_assessment.to_dict(),
                    "barter_signal_assessment": (
                        barter_assessment.to_dict()
                    ),
                    "compatibility_assessment": compatibility.to_dict(),
                    "live_campaign_assessment": (
                        live_assessment.to_dict()
                        if live_assessment is not None
                        else None
                    ),
                    "barter_review_threshold": {
                        "q1": phase_a_q1,
                        "q3": phase_a_q3,
                        "formula": "q3 + 3 * (q3 - q1)",
                        "threshold": barter_review_threshold,
                        "review_required": high_audience,
                        "preferred_review_threshold": (
                            config.eligibility.preferred_followers_max
                        ),
                        "preferred_range_review_required": (
                            live_assessment.audience_review_required
                            if live_assessment is not None
                            else False
                        ),
                    },
                    "discovery_confidence": discovery_confidence.score,
                    "score_preview": (
                        score_preview.to_dict()
                        if score_preview is not None
                        else None
                    ),
                    "eligibility": decision.to_dict(),
                }
            )
            if not decision.eligible or offer is None:
                excluded_records.append(
                    _excluded_profile(profile, decision)
                )
                continue

            score = score_preview
            if score is None:
                raise ValueError(
                    "eligible candidate is missing a reproducible score"
                )
            confidence_evidence = SignalEvidence(
                signal_type="discovery_confidence",
                source_field="provider_and_discovery",
                source_reference=profile.identity.username,
                evidence_text=discovery_confidence.explanation,
                observation_type="derived",
                url=profile.identity.profile_url,
            )
            all_evidence = _unique_evidence(
                (
                    *flatten_evidence(evidence),
                    *account_assessment.evidence,
                    *barter_assessment.barter_evidence,
                    *barter_assessment.no_barter_evidence,
                    *compatibility.evidence,
                    confidence_evidence,
                    *offer.evidence,
                )
            )
            collected_at = profile.collected_at or manifest.started_at
            audience_review_required = bool(
                live_assessment is not None
                and live_assessment.audience_review_required
            )
            eligible_results.append(
                CandidateResult(
                    platform=profile.identity.platform,
                    username=profile.identity.username,
                    profile_url=profile.identity.canonical_profile_url,
                    followers=metrics.followers,
                    median_likes=metrics.median_likes,
                    median_comments=metrics.median_comments,
                    engagement_rate=metrics.engagement_rate,
                    usable_posts=metrics.usable_posts,
                    sampled_posts=metrics.sampled_posts,
                    data_completeness=metrics.data_completeness,
                    short_video_share=metrics.short_video_share,
                    last_post_date=metrics.last_post_date,
                    score=score.score,
                    score_components=score.components,
                    selection_explanation=score.explanation,
                    evidence=all_evidence,
                    recent_post_url=offer.recent_post_url,
                    barter_offer=offer.text,
                    manual_verification_status=offer.manual_verification_status,
                    verification_notes=(
                        "Mandatory manual review before any use; no outreach "
                        "was sent."
                    ),
                    discovery_confidence=discovery_confidence.score,
                    eligibility_status=decision.status,
                    eligibility_reasons=decision.reasons,
                    query_ids=profile.query_ids,
                    provider=profile.provider,
                    collected_at=collected_at,
                    offer_generation_mode=offer.generation_mode,
                    source_exclusion_check=(
                        "clear: no Phase A username, URL, historical alias, "
                        "replacement, rejected false lead, Nike, or Apple match"
                    ),
                    account_type=account_assessment.account_type.value,
                    account_type_explanation=account_assessment.explanation,
                    barter_feasibility_review_required=(
                        audience_review_required
                    ),
                    barter_feasibility_explanation=(
                        (
                            "Audience is above the preferred automatic barter "
                            "range ceiling "
                            f"{config.eligibility.preferred_followers_max:,} "
                            "but does not exceed the frozen Phase A "
                            "outer-fence threshold "
                            f"{barter_review_threshold:,.1f}; manual barter "
                            "feasibility review is required."
                        )
                        if audience_review_required
                        else (
                            "Audience is within the configured preferred "
                            "automatic barter review range and does not exceed "
                            "the frozen Phase A outer-fence threshold "
                            f"{barter_review_threshold:,.1f}."
                        )
                    ),
                    content_themes=(
                        account_assessment.relevant_dimensions
                    ),
                    known_format_posts=known_format_posts(profile),
                    detected_content_language=(
                        compatibility.detected_content_language
                    ),
                    campaign_language_compatible=(
                        compatibility.campaign_language_compatible
                    ),
                    detected_geography=(
                        compatibility.detected_geography
                    ),
                    delivery_market_review_required=(
                        compatibility.delivery_market_review_required
                    ),
                    compatibility_explanation=(
                        compatibility.explanation
                    ),
                    barter_evidence=(
                        barter_assessment.barter_evidence
                    ),
                    no_barter_evidence=(
                        barter_assessment.no_barter_evidence
                    ),
                    campaign_bucket=(
                        (
                            "needs_manual_review"
                            if audience_review_required
                            else "barter_ready"
                        )
                        if config.mode == "live"
                        else ""
                    ),
                    campaign_status_reasons=(
                        (
                            "audience_above_100k_manual_barter_review",
                        )
                        if audience_review_required
                        else ()
                    ),
                )
            )

        ranked = rank_candidates(eligible_results)
        ranked = [
            replace(
                candidate,
                selection_explanation=selection_explanation(
                    candidate, rank, len(ranked)
                ),
            )
            for rank, candidate in enumerate(ranked, start=1)
        ]
        manifest.provider_run_ids = sorted(provider_run_ids)
        _sync_provider_usage(manifest, provider)
        source_exclusion_count = sum(
            item.get("exclusion_reason")
            in {"source_exclusion", "previous_live_run_exclusion"}
            for item in deduplicated.excluded_records
        )
        previous_live_run_exclusion_count = sum(
            item.get("exclusion_reason")
            == "previous_live_run_exclusion"
            for item in deduplicated.excluded_records
        )
        malformed_count = sum(
            item.get("exclusion_reason") == "malformed_discovery_record"
            for item in deduplicated.excluded_records
        )
        manifest.counts = {
            "generated_queries": len(queries),
            "raw_discovery_hits": len(discovery_hits),
            "source_exclusions": source_exclusion_count,
            "malformed_discovery_records": malformed_count,
            "duplicate_discoveries": int(
                deduplicated.report["duplicate_hit_count"]
            ),
            "unique_candidates": len(deduplicated.identities),
            "enriched_candidates": len(profiles),
            "eligible_candidates": len(ranked),
            "ineligible_candidates": len(profiles) - len(ranked),
            "excluded_records_total": len(excluded_records),
            "selected_candidates": 0,
        }
        if config.run_level_exclusions:
            manifest.counts["previous_live_run_exclusions"] = (
                previous_live_run_exclusion_count
            )
        # Persistent provider-response caching is intentionally outside this
        # bounded MVP; report that explicitly instead of inventing cache misses.
        manifest.cache = {"enabled": 0, "hits": 0, "misses": 0}
        try:
            selection = select_top_candidates(
                ranked,
                final_count=config.discovery.final_count,
                minimum_count=config.discovery.minimum_final_count,
            )
        except InsufficientCandidatePoolError as exc:
            near_misses = build_near_miss_candidates(enriched_records)
            manifest.counts["near_miss_candidates"] = len(near_misses)
            _record_failed_manifest(run_dir, manifest, exc)
            discovery_payload = [
                hit.to_dict(include_raw=True) for hit in discovery_hits
            ]
            try:
                generated_paths = generate_phase_b_failure_artifacts(
                    run_dir,
                    manifest=manifest,
                    queries=queries,
                    discovery_pool=discovery_payload,
                    deduplication_report=deduplicated.report,
                    excluded_candidates=excluded_records,
                    enriched_candidates=enriched_records,
                    eligible_candidates=ranked,
                    near_miss_candidates=near_misses,
                )
            except (OSError, ValueError, KeyError) as output_exc:
                raise OutputWriteError(
                    "Could not persist insufficient-pool audit artifacts: "
                    f"{output_exc}"
                ) from output_exc
            manifest.artifacts = {
                path.name: str(path) for path in generated_paths
            }
            _write_manifest(run_dir / "run_manifest.json", manifest)
            raise
        manifest.counts["selected_candidates"] = len(selection.selected)
        if selection.warning:
            manifest.warnings.append(selection.warning)
        if config.google_sheets.enabled and resolved_sheets_client is None:
            manifest.warnings.append(
                "Google Sheets output was configured but no authenticated "
                "client was injected; local artifacts were generated."
            )
        manifest.status = "completed"
        manifest.completed_at = datetime.now(timezone.utc)
        manifest.artifacts = {
            name: str(run_dir / name) for name in ARTIFACT_FILENAMES
        }

        discovery_payload = [
            hit.to_dict(include_raw=True) for hit in discovery_hits
        ]
        try:
            generated_paths = generate_phase_b_artifacts(
                run_dir,
                manifest=manifest,
                queries=queries,
                discovery_pool=discovery_payload,
                deduplication_report=deduplicated.report,
                excluded_candidates=excluded_records,
                enriched_candidates=enriched_records,
                eligible_candidates=ranked,
                selected_candidates=selection.selected,
                source_workbook=config.inputs.workbook,
                workbook_sheet=config.outputs.workbook_sheet,
            )
        except (OSError, ValueError, KeyError) as exc:
            raise OutputWriteError(
                f"Could not generate Phase B artifacts: {exc}"
            ) from exc

        sheets_receipt: GoogleSheetsWriteReceipt | None = None
        if config.google_sheets.enabled:
            if resolved_sheets_client is not None:
                try:
                    adapter = GoogleSheetsAdapter(
                        resolved_sheets_client,
                        config.google_sheets.spreadsheet_id or "",
                        worksheet_name=config.google_sheets.worksheet_name,
                    )
                    sheets_receipt = adapter.upsert(selection.selected)
                except Exception as exc:
                    raise OutputWriteError(
                        f"Google Sheets upsert failed: {exc}"
                    ) from exc
                manifest.artifacts["google_sheets"] = (
                    f"{sheets_receipt.spreadsheet_id}#"
                    f"{sheets_receipt.worksheet_name}"
                )
                _write_manifest(run_dir / "run_manifest.json", manifest)

        return PhaseBRunResult(
            manifest=manifest,
            run_dir=run_dir,
            generated_paths=tuple(generated_paths),
            selected_candidates=selection.selected,
            eligible_candidates=tuple(ranked),
            google_sheets_receipt=sheets_receipt,
        )
    except PhaseBError as exc:
        _sync_provider_usage(manifest, provider)
        _record_failed_manifest(run_dir, manifest, exc)
        raise
    except (OSError, ValueError, KeyError, csv.Error, json.JSONDecodeError) as exc:
        wrapped = InputValidationError(f"Phase B pipeline failed: {exc}")
        _sync_provider_usage(manifest, provider)
        _record_failed_manifest(run_dir, manifest, wrapped)
        raise wrapped from exc


def _sync_provider_usage(
    manifest: PhaseBRunManifest,
    provider: InstagramProvider | None,
) -> None:
    if not isinstance(provider, ApifyInstagramProvider):
        return
    manifest.provider_requests_made = provider.provider_requests_made
    manifest.budget_spent_usd = provider.budget_spent_usd
    manifest.provider_run_ids = list(
        dict.fromkeys(
            (*manifest.provider_run_ids, *provider.provider_run_ids)
        )
    )
    manifest.review_summary["provider_charges"] = [
        dict(item) for item in provider.actor_run_usage
    ]
    manifest.review_summary["provider_budget_guard"] = {
        "per_actor_run_cap_usd": (
            provider.config.max_total_charge_usd
        ),
        "combined_cap_usd": (
            provider.config.max_combined_charge_usd
        ),
        "at_most_one_search_run": True,
        "at_most_one_profile_run": True,
    }


def _build_provider(config: PhaseBConfig) -> InstagramProvider:
    if config.provider.type == "fixture":
        fixture = config.provider.fixture
        if fixture is None:
            raise ConfigurationError("fixture provider configuration is missing")
        return FixtureInstagramProvider(
            fixture.discovery_path, fixture.profiles_path
        )
    apify = config.provider.apify
    if apify is None:
        raise ConfigurationError("Apify provider configuration is missing")
    secrets = load_env_secrets(config.env_file)
    return ApifyInstagramProvider(
        apify,
        require_secret(secrets, "APIFY_TOKEN"),
    )


def _excluded_hit(
    hit: DiscoveryHit,
    reason: str,
    reasons: Iterable[str],
    *,
    exclusion_match: ExclusionMatch | None = None,
) -> dict[str, Any]:
    normalized = normalize_username(hit.username) or normalize_username(
        hit.profile_url
    )
    result: dict[str, Any] = {
        "platform": hit.platform,
        "username": hit.username,
        "profile_url": hit.profile_url,
        "normalized_username": normalized or "",
        "exclusion_reason": reason,
        "exclusion_reasons": list(reasons),
        "query_ids": list(hit.query_ids),
        "provider": hit.provider,
    }
    if exclusion_match is not None:
        result["exclusion_registry_match"] = exclusion_match.to_dict()
    return result


def _excluded_profile(
    profile: CreatorProfile, decision: EligibilityDecision
) -> dict[str, Any]:
    return {
        "platform": profile.identity.platform,
        "username": profile.identity.username,
        "profile_url": profile.identity.canonical_profile_url,
        "normalized_username": profile.identity.normalized_username,
        "exclusion_reason": decision.status.value,
        "exclusion_reasons": list(decision.reasons),
        "query_ids": list(profile.query_ids),
        "provider": profile.provider,
    }


def _unique_evidence(
    evidence: Iterable[SignalEvidence],
) -> tuple[SignalEvidence, ...]:
    unique: list[SignalEvidence] = []
    seen: set[tuple[str, str, str, str, str | None]] = set()
    for item in evidence:
        key = (
            item.signal_type,
            item.source_field,
            item.source_reference,
            item.evidence_text,
            item.url,
        )
        if key not in seen:
            seen.add(key)
            unique.append(item)
    return tuple(unique)


def _read_json_object(path: Path, label: str) -> Mapping[str, Any]:
    try:
        value = json.loads(Path(path).read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise InputValidationError(f"{label} file not found: {path}") from exc
    except json.JSONDecodeError as exc:
        raise InputValidationError(
            f"{path}: invalid JSON at line {exc.lineno}, column {exc.colno}"
        ) from exc
    if not isinstance(value, Mapping):
        raise InputValidationError(f"{path}: expected a JSON object")
    return value


def _validate_source_analysis(path: Path) -> None:
    required = {"username", "status", "engagement_rate_pct"}
    try:
        with Path(path).open("r", encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            missing = required - set(reader.fieldnames or ())
            if missing:
                raise InputValidationError(
                    f"{path}: missing columns: {', '.join(sorted(missing))}"
                )
            rows = list(reader)
    except FileNotFoundError as exc:
        raise InputValidationError(f"source analysis not found: {path}") from exc
    except csv.Error as exc:
        raise InputValidationError(f"{path}: malformed CSV: {exc}") from exc
    if not any(
        row.get("status") == "creator"
        and str(row.get("engagement_rate_pct") or "").strip()
        for row in rows
    ):
        raise InputValidationError(
            f"{path}: no scored Phase A creator engagement reference values"
        )


def _validate_workbook(path: Path) -> None:
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
        workbook.close()
    except (OSError, ValueError) as exc:
        raise InputValidationError(f"{path}: invalid XLSX workbook: {exc}") from exc


def _record_failed_manifest(
    run_dir: Path, manifest: PhaseBRunManifest, error: PhaseBError
) -> None:
    manifest.status = "failed"
    manifest.completed_at = datetime.now(timezone.utc)
    error_record = {
        "category": error.category,
        "message": str(error),
        "details": error.details,
    }
    if error_record not in manifest.errors:
        manifest.errors.append(error_record)
    run_dir.mkdir(parents=True, exist_ok=True)
    try:
        _write_manifest(run_dir / "run_manifest.json", manifest)
    except OSError:
        # Preserve the original, more actionable pipeline error.
        pass


def _write_manifest(path: Path, manifest: PhaseBRunManifest) -> None:
    path.write_text(
        json.dumps(
            manifest.to_dict(),
            ensure_ascii=False,
            indent=2,
            sort_keys=True,
        )
        + "\n",
        encoding="utf-8",
    )


__all__ = [
    "DeduplicationOutcome",
    "PhaseBRunResult",
    "deduplicate_discovery_hits",
    "run_phase_b",
    "validate_phase_b_config",
]
